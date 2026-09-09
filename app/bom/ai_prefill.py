"""Tier 3 of the ingestion ladder: AI-assisted *template pre-fill*.

What this is NOT: a path from a document into validation. The plan is firm
that model output never enters the rules engine directly. This module turns
an unrecognised quote (PDF, DOCX, odd spreadsheet, plain text) into a filled
copy of our strict xlsx template; the user downloads it, reviews it, and
uploads it like any other BOM — where the same strict template parser
(bom.parsers.template) is the only thing that produces a NormalizedBOM.

Gating: the feature exists only when ANTHROPIC_API_KEY is set (and the
``anthropic`` package is importable). Without a key the UI hides the panel
and the route answers 503; the blank template is always available.

The call is one Messages request with a JSON-schema constrained output, so
the response is valid JSON by construction; it is still sanitised through
NormalizedBOM.from_dict (unknown categories → 'other', bad quantities → 1).
PDFs are sent as a document block — no PDF text extraction dependency.
"""
import base64
import csv
import io
import json
import os
from typing import Any, Dict, Optional, Tuple

from bom.normalize import CATEGORIES, NormalizedBOM, VENDORS

DEFAULT_MODEL = "claude-opus-5"
MAX_TEXT_CHARS = 60000
MAX_FILE_BYTES = 10 * 1024 * 1024

PREFILL_EXTENSIONS = (".xlsx", ".csv", ".pdf", ".docx", ".txt")


class PrefillError(Exception):
    """User-presentable failure (unsupported file, model declined, bad JSON)."""


class PrefillUnavailable(PrefillError):
    """No API key / SDK: the feature is switched off on this deployment."""


SYSTEM_PROMPT = """You are a hardware BOM (bill of materials) parser for a data-centre hardware compatibility tool.

Extract the hardware line items from the quote or BOM the user provides. Vendors are typically Dell, Lenovo, Supermicro or HPE, in many layouts (Dell quotes and service-tag exports, Lenovo DCSC configurator exports, distributor bids, hand-written lists).

Return one object per distinct server configuration (a quote often has a production build and a DR build). For each configuration give:
- name: the configuration's name as written, else "Config 1", "Config 2", ...
- serverModel: the server model (e.g. "PowerEdge R760", "ThinkSystem SR650 V3"), or null.
- nodeCount: how many servers of this configuration the quote contains (the quantity on the server/chassis line), or null if unclear.
- components: one entry per line item with partNumber (the vendor SKU or feature code exactly as printed, else null), description (verbatim, without trailing part numbers), quantity (the TOTAL quantity across all servers of the configuration, exactly as printed) and category.

Categories:
- "controller": HBA, RAID card, PERC, storage controller
- "boss": BOSS card or any M.2 RAID/mirroring boot adapter
- "storage": hard drives, SSDs, NVMe drives (keep the interface words NVMe/SAS/SATA/HDD/SSD in the description)
- "nic": network adapters, OCP/PCIe NICs, LOM, rNDC
- "cpu": processors
- "memory": DIMMs / RDIMMs
- "chassis": server chassis, backplane, the server base line itself
- "gpu": GPU cards
- "other": power supplies, rails, cables, fans, TPM, bezels, licences, services, labels, documentation, settings

Rules: keep "No BOSS", "No Controller", "No HBA", "No RAID", "BOSS Blank", "LOM Blank" as "other" (they record a deliberate absence). Skip pure info/label/shipping/firmware lines (INFO, LBL, SRV,SW, SHP MTL, GDE, PREP MTL). Do not invent parts that are not in the document. Do not sum or split quantities."""

OUTPUT_SCHEMA = {
    "type": "object",
    "properties": {
        "vendor": {"type": "string", "enum": list(VENDORS)},
        "configs": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "serverModel": {"type": ["string", "null"]},
                    "nodeCount": {"type": ["integer", "null"]},
                    "components": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "partNumber": {"type": ["string", "null"]},
                                "description": {"type": "string"},
                                "quantity": {"type": "integer"},
                                "category": {"type": "string", "enum": list(CATEGORIES)},
                            },
                            "required": ["partNumber", "description", "quantity", "category"],
                            "additionalProperties": False,
                        },
                    },
                },
                "required": ["name", "serverModel", "nodeCount", "components"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["vendor", "configs"],
    "additionalProperties": False,
}


def api_key() -> Optional[str]:
    return os.environ.get("ANTHROPIC_API_KEY") or None


def model_name() -> str:
    return os.environ.get("BOM_PREFILL_MODEL") or DEFAULT_MODEL


def sdk_available() -> bool:
    try:
        import anthropic  # noqa: F401
        return True
    except ImportError:
        return False


def available() -> bool:
    return bool(api_key()) and sdk_available()


# ── document extraction ─────────────────────────────────────────────────────

def _xlsx_text(path: str) -> str:
    from openpyxl import load_workbook
    from xlsx_utils import MAX_SHEET_COLS, MAX_SHEET_ROWS
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:20]:
        parts.append("=== Sheet: %s ===" % ws.title)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                break
            cells = ["" if v is None else str(v).strip() for v in row[:MAX_SHEET_COLS]]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _csv_text(path: str) -> str:
    with open(path, "rb") as fh:
        raw = fh.read()
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise PrefillError("The CSV file is not readable text.")
    rows = list(csv.reader(io.StringIO(text)))
    return "\n".join("\t".join(c.strip() for c in r) for r in rows if any(c.strip() for c in r))


def _docx_text(path: str) -> str:
    import docx  # python-docx, already a dependency
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def extract_document(path: str, filename: str) -> Tuple[str, Any]:
    """('text', str) or ('pdf', base64 str) for the user message."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in PREFILL_EXTENSIONS:
        raise PrefillError("Unsupported file type %s. Accepted: %s" % (ext or "(none)", ", ".join(PREFILL_EXTENSIONS)))
    if os.path.getsize(path) > MAX_FILE_BYTES:
        raise PrefillError("File too large (max 10 MB).")
    if ext == ".pdf":
        with open(path, "rb") as fh:
            head = fh.read(5)
            fh.seek(0)
            if head[:4] != b"%PDF":
                raise PrefillError("The file does not look like a PDF.")
            return "pdf", base64.b64encode(fh.read()).decode("ascii")
    if ext == ".xlsx":
        text = _xlsx_text(path)
    elif ext == ".csv":
        text = _csv_text(path)
    elif ext == ".docx":
        text = _docx_text(path)
    else:
        with open(path, "rb") as fh:
            text = fh.read().decode("utf-8", "replace")
    if not text.strip():
        raise PrefillError("The document contains no readable text.")
    if len(text) > MAX_TEXT_CHARS:
        raise PrefillError("The document is too long to pre-fill automatically (%d characters, max %d). "
                           "Fill the blank template instead." % (len(text), MAX_TEXT_CHARS))
    return "text", text


def _user_content(kind: str, payload: Any, filename: str):
    safe_name = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in (filename or "bom"))[:100]
    if kind == "pdf":
        return [
            {"type": "document",
             "source": {"type": "base64", "media_type": "application/pdf", "data": payload}},
            {"type": "text", "text": "Extract the hardware BOM from this quote (filename: %s)." % safe_name},
        ]
    return [{"type": "text",
             "text": "BOM document (filename: %s):\n```\n%s\n```" % (safe_name, payload)}]


def _client():
    if not api_key():
        raise PrefillUnavailable("AI pre-fill is not configured on this server (no API key).")
    try:
        import anthropic
    except ImportError:
        raise PrefillUnavailable("AI pre-fill is not available: the anthropic package is not installed.")
    return anthropic.Anthropic(api_key=api_key(), max_retries=2)


def extract_bom(path: str, filename: str, client=None, model: Optional[str] = None) -> NormalizedBOM:
    """One model call → sanitised NormalizedBOM. ``client`` is injectable for tests."""
    kind, payload = extract_document(path, filename)
    client = client or _client()
    response = client.messages.create(
        model=model or model_name(),
        max_tokens=16000,
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": _user_content(kind, payload, filename)}],
        output_config={"format": {"type": "json_schema", "schema": OUTPUT_SCHEMA}},
    )
    if getattr(response, "stop_reason", None) == "refusal":
        raise PrefillError("The model declined to process this document.")
    if getattr(response, "stop_reason", None) == "max_tokens":
        raise PrefillError("The document is too long to pre-fill in one pass; fill the blank template instead.")
    text = None
    for block in getattr(response, "content", None) or []:
        if getattr(block, "type", None) == "text":
            text = block.text
            break
    if not text:
        raise PrefillError("The model returned no content.")
    try:
        data = json.loads(text)
    except ValueError:
        raise PrefillError("The model returned malformed JSON.")
    bom = NormalizedBOM.from_dict(data)
    if not bom.configs:
        raise PrefillError("No hardware configurations were found in the document.")
    return bom


def prefill_template(path: str, filename: str, client=None, model: Optional[str] = None,
                     lang: str = "en") -> bytes:
    """The filled xlsx template for ``path``. The caller streams it back to
    the user; nothing is stored."""
    from bom.parsers.template import build_template_bytes
    bom = extract_bom(path, filename, client=client, model=model)
    return build_template_bytes(bom=bom, lang=lang)


def capabilities() -> Dict[str, Any]:
    return {
        "ai_prefill_available": available(),
        "prefill_extensions": list(PREFILL_EXTENSIONS),
        "prefill_model": model_name() if available() else None,
    }
