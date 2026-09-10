"""Tier 2 of the ingestion ladder: our own strict xlsx template.

Why strict: the template exists so that *any* BOM — a PDF, an email, an AI
pre-fill — can reach the validator through a file whose meaning is not in
doubt. Leniency would reintroduce exactly the guessing the plan forbids, so
parse_template() collects every problem with its row number and raises them
together (TemplateError) instead of silently defaulting: the user fixes the
sheet in one pass and re-uploads.

Layout (build spec §9, plus one optional column):

  Config | Server model | Vendor | Part number | Description | Quantity | Category | Nodes

  - Quantities are TOTALS across all machines of the config, like every
    deterministic parser; 'Nodes' (optional) is the machine count of the
    config, read from the first row of the config that fills it, so
    node_count round-trips through the template.
  - Vendor and Category are data-validation lists (Excel shows a dropdown);
    the parser re-checks them case-insensitively because validation lists do
    not survive every spreadsheet tool.
  - The workbook is marked twice — defined name SC_BOM_TEMPLATE = BOM!$A$1
    and a cell comment 'sc-bom-template v1' on A1 — so detect_format can
    recognise it before trying any vendor shape.

Column widths and the header styling mirror admin_routes.catalog_template so
the two downloadable templates look like one product.
"""
import io
from typing import Dict, List, Optional

from bom.normalize import CATEGORIES, VENDORS, BOMComponent, BOMConfig, NormalizedBOM
from bom.parsers.common import cell, head_rows, raw_text_from_rows, s, sheet_matrix, to_int

FORMAT = 'template'

SHEET = 'BOM'
INSTRUCTIONS_SHEET = 'Instructions'
DEFINED_NAME = 'SC_BOM_TEMPLATE'
MARKER_COMMENT = 'sc-bom-template v1'

COLUMNS = ['Config', 'Server model', 'Vendor', 'Part number', 'Description',
           'Quantity', 'Category']
OPTIONAL_COLUMNS = ['Nodes']
ALL_COLUMNS = COLUMNS + OPTIONAL_COLUMNS
VALIDATION_ROWS = 500
_WIDTHS = [22, 24, 12, 20, 60, 10, 12, 8]

_CATEGORY_HELP = [
    ('cpu', 'Processors (Intel Xeon, AMD EPYC).'),
    ('memory', 'RAM modules (RDIMM, UDIMM, LRDIMM, DIMM).'),
    ('storage', 'Data drives: HDDs, SSDs, NVMe (not M.2 boot sticks - put those under other).'),
    ('controller', 'HBA / RAID cards (HBA355i, PERC, 440-16i, 9350-8i, AOC-S3008).'),
    ('nic', 'Network adapters: OCP / PCIe NICs, rNDC, on-board LOM.'),
    ('chassis', 'The server / chassis / backplane line (its quantity is the number of nodes).'),
    ('boss', 'BOSS / M.2 RAID boot cards.'),
    ('gpu', 'GPUs and accelerators.'),
    ('other', 'Everything else: PSU, fans, rails, cables, TPM, bezel, licences, services,'
              ' and absence lines such as "No BOSS Card" or "LOM Blank".'),
]

_INSTRUCTIONS = {
    'en': [
        ('How to use this template', None),
        ('1.', 'Download this file, fill the BOM sheet, upload it on the project page ("Check a BOM").'),
        ('2.', 'One row per line item of the quote. Do not merge cells, do not rename the sheet'
               ' or the header row.'),
        ('3.', 'Quantities are TOTALS across all nodes of the config (3 nodes x 2 CPUs = 6).'),
        ('4.', 'Config: a name for each build (repeat it on every row of that build). Blank = "Config 1".'),
        ('5.', 'Nodes (optional): the number of machines in the config, on the first row of the config.'
               ' When blank, the quantity of the chassis/server line is used.'),
        ('6.', 'Server model: e.g. "PowerEdge R760", "ThinkSystem SR650 V4", "SYS-511R-M".'),
        ('7.', 'Vendor: Dell, Lenovo, Supermicro, HPE or Unknown.'),
        ('8.', 'Part number: the vendor part number / SKU / feature code when known; blank is allowed.'),
        ('9.', 'Description: the vendor description, verbatim. Required.'),
        ('10.', 'Category: one of the values below. Required.'),
        ('', None),
        ('Categories', None),
    ] + [(c, h) for c, h in _CATEGORY_HELP] + [
        ('', None),
        ('Keep absence lines', 'Rows such as "No BOSS Card", "No Controller", "No RAID" or "LOM Blank"'
                               ' are useful evidence - keep them with category "other".'),
    ],
}


class TemplateError(ValueError):
    """Raised by parse_template with every problem found, row-numbered."""

    def __init__(self, errors: List[str]):
        self.errors = list(errors)
        super().__init__('; '.join(self.errors))


# ─── build ────────────────────────────────────────────────────────────────────

def build_template_bytes(bom: Optional[NormalizedBOM] = None, lang: str = 'en') -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    header_font = Font(bold=True, color='FFFFFF', size=10)
    header_fill = PatternFill('solid', fgColor='003A70')
    header_align = Alignment(horizontal='center', wrap_text=True)

    ws.append(ALL_COLUMNS)
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
    for i, width in enumerate(_WIDTHS):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = width
    ws.freeze_panes = 'A2'
    ws['A1'].comment = Comment(MARKER_COMMENT, 'Scale Computing sizer')

    dv_vendor = DataValidation(type='list', formula1='"%s"' % ','.join(VENDORS),
                               allow_blank=True, showErrorMessage=True,
                               errorTitle='Vendor', error='Pick a vendor from the list.')
    dv_cat = DataValidation(type='list', formula1='"%s"' % ','.join(CATEGORIES),
                            allow_blank=False, showErrorMessage=True,
                            errorTitle='Category', error='Pick a category from the list.')
    ws.add_data_validation(dv_vendor)
    ws.add_data_validation(dv_cat)
    dv_vendor.add('C2:C%d' % VALIDATION_ROWS)
    dv_cat.add('G2:G%d' % VALIDATION_ROWS)

    if bom is not None:
        for config in bom.configs:
            first = True
            for comp in config.components:
                ws.append([
                    config.name,
                    config.server_model or '',
                    bom.vendor or 'Unknown',
                    comp.part_number or '',
                    comp.description,
                    int(comp.quantity),
                    comp.category,
                    (config.node_count if first and config.node_count else None),
                ])
                first = False

    # Marker for detect_format: a defined name survives re-saves in every
    # spreadsheet tool we have seen; the A1 comment is for humans.
    dn = DefinedName(DEFINED_NAME, attr_text="'%s'!$A$1" % SHEET)
    wb.defined_names[DEFINED_NAME] = dn

    ws_help = wb.create_sheet(INSTRUCTIONS_SHEET)
    ws_help.column_dimensions['A'].width = 22
    ws_help.column_dimensions['B'].width = 110
    for key, text in _INSTRUCTIONS.get(lang, _INSTRUCTIONS['en']):
        ws_help.append([key, text])
        if text is None and key:
            ws_help.cell(row=ws_help.max_row, column=1).font = Font(bold=True, size=12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── detect / parse ───────────────────────────────────────────────────────────

def _header_ok(rows) -> bool:
    got = [cell(rows, 1, c).lower() for c in range(1, len(COLUMNS) + 1)]
    return got == [h.lower() for h in COLUMNS]


def detect(wb) -> bool:
    try:
        names = set(wb.defined_names.keys())
    except AttributeError:                      # older openpyxl API
        names = {d.name for d in wb.defined_names.definedName}
    if DEFINED_NAME in names:
        return True
    if SHEET in wb.sheetnames:
        return _header_ok(head_rows(wb[SHEET], 1, len(ALL_COLUMNS)))
    return False


def _canon(value: str, choices) -> Optional[str]:
    low = value.strip().lower()
    for choice in choices:
        if choice.lower() == low:
            return choice
    return None


def _positive_int(value, label: str, row: int, errors: List[str], required: bool) -> Optional[int]:
    text = s(value)
    if not text:
        if required:
            errors.append('Row %d: %s is required.' % (row, label))
        return None
    n = to_int(text, default=0)
    try:
        exact = float(text.replace(',', '.'))
    except ValueError:
        exact = None
    if n <= 0 or exact is None or exact != n:
        errors.append('Row %d: %s must be a positive whole number (got %r).' % (row, label, text))
        return None
    return n


def parse_template(path: str) -> NormalizedBOM:
    from bom.parsers import UnrecognizedFormat
    from bom.parsers.common import load_workbook_safe
    wb = load_workbook_safe(path)
    try:
        if SHEET not in wb.sheetnames:
            raise TemplateError(["Sheet '%s' is missing." % SHEET])
        rows = sheet_matrix(wb[SHEET])
    finally:
        wb.close()
    if not rows or not _header_ok(rows):
        raise TemplateError([
            'Row 1: header must be exactly "%s"%s.' % (
                ' | '.join(COLUMNS), ' (optional 8th column "Nodes")')])
    has_nodes = cell(rows, 1, 8).lower() == 'nodes' if len(rows[0]) >= 8 else False
    if len(rows[0]) >= 8 and cell(rows, 1, 8) and not has_nodes:
        raise TemplateError(['Row 1: the 8th column may only be "Nodes" (got %r).' % cell(rows, 1, 8)])

    errors: List[str] = []
    configs: Dict[str, BOMConfig] = {}
    order: List[str] = []
    vendor: Optional[str] = None
    for r in range(2, len(rows) + 1):
        row = rows[r - 1]
        if all(s(v) == '' for v in row):
            continue
        name = cell(rows, r, 1) or 'Config 1'
        model = cell(rows, r, 2) or None
        vendor_text = cell(rows, r, 3)
        part = cell(rows, r, 4) or None
        desc = cell(rows, r, 5)
        qty = _positive_int(row[5] if len(row) > 5 else None, 'Quantity', r, errors, required=True)
        cat_text = cell(rows, r, 7)
        nodes = _positive_int(row[7] if len(row) > 7 else None, 'Nodes', r, errors,
                              required=False) if has_nodes else None

        if not desc:
            errors.append('Row %d: Description is required.' % r)
        category = _canon(cat_text, CATEGORIES) if cat_text else None
        if category is None:
            errors.append('Row %d: Category must be one of %s (got %r).'
                          % (r, ', '.join(CATEGORIES), cat_text))
        if vendor_text:
            v = _canon(vendor_text, VENDORS)
            if v is None:
                errors.append('Row %d: Vendor must be one of %s (got %r).'
                              % (r, ', '.join(VENDORS), vendor_text))
            elif vendor is None and v != 'Unknown':
                vendor = v
        if errors and (qty is None or not desc or category is None):
            continue
        config = configs.get(name)
        if config is None:
            config = BOMConfig(name=name, server_model=model, components=[], node_count=None)
            configs[name] = config
            order.append(name)
        elif config.server_model is None and model:
            config.server_model = model
        if nodes is not None:
            if config.node_count is None:
                config.node_count = nodes
            elif config.node_count != nodes:
                errors.append('Row %d: Nodes for config %r conflicts with an earlier row (%d vs %d).'
                              % (r, name, config.node_count, nodes))
        config.components.append(BOMComponent(part_number=part, description=desc,
                                              quantity=qty, category=category))
    if errors:
        raise TemplateError(errors)
    if not order:
        raise TemplateError(['The BOM sheet has no line items.'])
    return NormalizedBOM(vendor=vendor or 'Unknown', configs=[configs[k] for k in order],
                         raw_text=raw_text_from_rows(rows))
