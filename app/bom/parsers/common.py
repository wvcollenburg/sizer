"""Helpers every deterministic BOM parser shares.

Why this module exists: the eight file shapes we recognise (Lenovo DCSC,
Dell service-tag export, Dell quote, D&H bid, Dell VNET module export and
three hand-made Dell lists) all end up asking the same three questions of a
line item — "is this junk?", "is this an absence indicator the rules engine
needs?" and "which of the nine categories is it?" — and they all have to
coerce the same kinds of sloppy cells (floats for integers, '3,0', NBSP).
Keeping the answers here means a category fix lands in every parser at once
and the archive-optional parity test (tests/test_bom_parsers.py) exercises
one implementation.

Design rules that matter downstream:

  - categorize() reproduces the categories SC//Design's LLM extractor gave
    the 26 archived BOMs (survey §4), with the survey's documented fixes:
    spec-only lines ('6400MT/s RDIMMs') are dropped, 'GB Base-T' (with a
    space) is a NIC, and a motherboard *with an on-board LOM* is a NIC so
    the LOM/backplane-over-VLAN info finding still fires.
  - Absence indicators ('No BOSS', 'BOSS Blank', 'No Controller', 'No RAID',
    'LOM Blank', 'No Additional Processor'...) are NEVER dropped: bom.rules
    treats them as evidence that a component was deliberately left out
    (is_absence_indicator), so they must survive as category 'other'.
  - Everything else that is pure order-entry noise (labels, fillers, shipping
    material, documentation, passwords, BIOS/iDRAC settings, support SKUs)
    is dropped before categorisation so it can never be mistaken for
    hardware by a keyword hit ('Power Saving Dell Active Power Controller'
    is not a controller).
  - Descriptions stay exactly as the source wrote them (after the Dell
    30-character de-wrap); we classify pt-BR DCSC exports by adding the
    Portuguese keywords rather than translating.

Pure Python + openpyxl; no Flask, no DB.
"""
import math
import re
import warnings
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple, Union

from xlsx_utils import MAX_SHEET_COLS, MAX_SHEET_ROWS, SheetTooLargeError
from bom.normalize import BOMComponent
# Quantity coercion lives with the normalized shape so the agent path and the
# sheet parsers agree on '3,0'; re-exported here for the parsers.
from bom.normalize import to_int  # noqa: F401

# ─── identifiers ─────────────────────────────────────────────────────────────
# One spelling per identifier, with the anchored / unanchored variants the
# parsers need named side by side. Two patterns that look alike but are not
# the same set are kept apart on purpose and say so.

_DELL_SKU_CORE = r'\d{3}-[A-Z0-9]{4}'                   # 405-AAZF, 892-9155
_LENOVO_CTO_CORE = r'7[0-9A-Z]{3}CTO\d+WW'              # 7DGDCTO1WW, 7D73CTO3WW
_LENOVO_BLOCK_CORE = r'[0-9A-Z]{4}CT[OS0-9][0-9A-Z]WW'  # machine (CTO), software and service (CTS) CTOs
_LENOVO_CM_CORE = r'\d{4}[A-Z]{2}\d'                    # 5374CM1 configuration instruction
_LENOVO_FC_CORE = r'[A-Z0-9]{4}'                        # BLKJ, C5R6 feature codes
_LENOVO_PART_CORE = r'[0-9][A-Z0-9]{9}'                 # 4X77A90963 part numbers

DELL_SKU = re.compile(r'^%s$' % _DELL_SKU_CORE)
# A SKU on a line of its own inside a text blob (detect_vendor).
DELL_SKU_LINE = re.compile(r'^%s$' % _DELL_SKU_CORE, re.M)
# Unanchored, plus the five-digit bundle form (486-85584). Wrapped SKU cells
# run together ('405-AAXY750-ACFR'), so the four-character form must not
# swallow the next SKU's first digit.
DELL_SKU_ANY = re.compile(r'\d{3}-(?:\d{5}(?![0-9])|[A-Z0-9]{4})')
# A SKU glued to the end of a word ('Server210-BDZY'): (text before, sku).
DELL_SKU_AT_END = re.compile(r'(.*?)(%s)$' % _DELL_SKU_CORE)

LENOVO_CTO = re.compile(r'^%s$' % _LENOVO_CTO_CORE)
LENOVO_CTO_WORD = re.compile(r'\b%s\b' % _LENOVO_CTO_CORE)
LENOVO_FC = re.compile(r'^%s$' % _LENOVO_FC_CORE)
LENOVO_PART = re.compile(r'^%s$' % _LENOVO_PART_CORE)
# Block starters of a DCSC list: machine / software / service CTOs.
LENOVO_BLOCK = re.compile(r'^%s$' % _LENOVO_BLOCK_CORE)
# The Arrow quote sheet also starts a block on a configuration instruction
# (5374CM1), whose children must not land in the machine's parts.
LENOVO_BLOCK_OR_CM = re.compile(r'^(?:%s|%s)$' % (_LENOVO_BLOCK_CORE, _LENOVO_CM_CORE))
# Any code the numbered DCSC list may carry in its part column.
LENOVO_LIST_CODE = re.compile(r'^(?:%s|%s|%s|%s)$' % (
    _LENOVO_FC_CORE, _LENOVO_BLOCK_CORE, _LENOVO_CM_CORE, _LENOVO_PART_CORE))
# The build order's feature-code column. NOT the same set as
# LENOVO_LIST_CODE: its CTO form is wider ('CT' + any two characters) and it
# has no ten-character part number — kept as written so the pattern-miss
# bookkeeping of that parser does not move.
LENOVO_BUILD_CODE = re.compile(r'^(?:%s|[0-9A-Z]{4}CT[A-Z0-9]{2}WW|%s)$' % (
    _LENOVO_FC_CORE, _LENOVO_CM_CORE))

# A leading "8x " / "4 X " multiplier on a hand-typed description.
NX_PREFIX = re.compile(r'^(\d+)\s*[xX]\s+(.*)$')


# ─── cell coercion ────────────────────────────────────────────────────────────

def s(v: Any) -> str:
    """Cell value as a stripped string. None -> ''. Integral floats print as
    ints ('3.0' cells are a common openpyxl artefact) so text comparisons
    such as `== 'Qty'` and part-number regexes behave."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        if math.isfinite(v) and v == int(v):
            return str(int(v))
        return str(v)
    return str(v).replace('\xa0', ' ').strip()


# ─── workbook access (bounded) ────────────────────────────────────────────────

# Decompression-bomb limits for the zip pre-check below: any single member's
# uncompressed size, all members together, and (for members big enough for it
# to matter) the deflate ratio.
MAX_XLSX_MEMBER_BYTES = 50 * 1024 * 1024
MAX_XLSX_TOTAL_BYTES = 200 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 100
_RATIO_CHECK_MIN_BYTES = 5 * 1024 * 1024


def check_zip_bomb(path: str) -> None:
    """Reject a decompression bomb BEFORE openpyxl touches the file.

    openpyxl (read_only included) parses xl/sharedStrings.xml in full into a
    Python list at load time, so the MAX_SHEET_ROWS/MAX_SHEET_COLS caps never
    see an oversized strings part: a sub-MB upload can inflate to gigabytes
    and OOM the worker. The zip central directory states every member's
    uncompressed size, so this is a cheap metadata check. Files that are not
    readable zips pass through — openpyxl raises its own error for those."""
    import zipfile
    try:
        zf = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError):
        return
    with zf:
        total = 0
        for info in zf.infolist():
            size = int(info.file_size or 0)
            total += size
            ratio = (size / info.compress_size) if info.compress_size else 0
            if (size > MAX_XLSX_MEMBER_BYTES
                    or (size > _RATIO_CHECK_MIN_BYTES
                        and ratio > MAX_XLSX_COMPRESSION_RATIO)):
                raise SheetTooLargeError(
                    "The workbook is too large to process safely: part %r "
                    "expands to %d MB." % (info.filename, size // (1024 * 1024)))
            if total > MAX_XLSX_TOTAL_BYTES:
                raise SheetTooLargeError(
                    "The workbook is too large to process safely: its parts "
                    "expand to more than %d MB in total."
                    % (MAX_XLSX_TOTAL_BYTES // (1024 * 1024)))


def load_workbook_safe(path: str):
    """read_only + data_only like every other importer in the app; openpyxl's
    'no default style' warning on some Lenovo exports is noise. The zip
    pre-check runs first — read_only mode still loads sharedStrings.xml
    eagerly, so the row/col caps alone cannot stop a decompression bomb."""
    check_zip_bomb(path)
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', module='openpyxl')
        return load_workbook(path, read_only=True, data_only=True)


XLS_MAGIC = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'    # OLE2 compound file (BIFF .xls)
MAX_XLS_SHEETS = 20


def is_xls(path: str) -> bool:
    try:
        with open(path, 'rb') as fh:
            return fh.read(8) == XLS_MAGIC
    except OSError:
        return False


def xls_to_xlsx(path: str) -> str:
    """Copy a legacy .xls workbook into a temporary .xlsx (caller unlinks).

    Every deterministic parser reads openpyxl workbooks; converting once lets
    an .xls export of a known layout go through the same detector ladder as
    its .xlsx twin instead of straight to the (paid) agent. Values only: the
    parsers never look at styles. Strings are stored as text so a cell
    reading '=...' stays literal (load_workbook_safe is data_only, and a
    formula without a cached value would come back empty). The row/column
    caps of the xlsx path apply; xlrd itself holds the whole file in memory,
    which the 10 MB upload cap bounds."""
    import os
    import tempfile
    import xlrd
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    if not is_xls(path):
        raise ValueError('not an .xls workbook')
    try:
        book = xlrd.open_workbook(path, on_demand=True)
    except Exception as exc:                  # xlrd raises several unrelated types
        raise ValueError('unreadable .xls workbook: %s' % exc)
    wb = Workbook()
    wb.remove(wb.active)
    try:
        for index in range(min(book.nsheets, MAX_XLS_SHEETS)):
            sheet = book.sheet_by_index(index)
            if sheet.nrows > MAX_SHEET_ROWS:
                raise SheetTooLargeError(
                    "Sheet '%s' exceeds the maximum of %d rows." % (sheet.name, MAX_SHEET_ROWS))
            ws = wb.create_sheet(title=(sheet.name or 'Sheet%d' % (index + 1))[:31])
            for r in range(sheet.nrows):
                for c in range(min(sheet.ncols, MAX_SHEET_COLS)):
                    cell_obj = sheet.cell(r, c)
                    if cell_obj.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell_obj.value
                    if cell_obj.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                        value = int(value)
                    elif cell_obj.ctype == xlrd.XL_CELL_ERROR:
                        continue
                    if isinstance(value, str):
                        # Control characters are legal in BIFF, not in xlsx.
                        value = ILLEGAL_CHARACTERS_RE.sub('', value)
                    target = ws.cell(row=r + 1, column=c + 1, value=value)
                    if isinstance(value, str):
                        target.data_type = 's'
            book.unload_sheet(index)
    finally:
        book.release_resources()
    if not wb.worksheets:
        wb.create_sheet('Sheet1')
    fd, out = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(out)
    return out


def sheet_matrix(ws, max_rows: int = MAX_SHEET_ROWS) -> List[Tuple]:
    """Materialise a worksheet as a list of row tuples, bounded by the same
    caps xlsx_utils applies to RVTools/LiveOptics uploads: an .xlsx is a ZIP
    and a declared used-range of millions of empty cells would otherwise
    expand until the worker OOMs. Trailing all-None rows are kept (row
    numbers must stay 1:1 with the sheet for error messages)."""
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True, max_col=MAX_SHEET_COLS)):
        if i >= max_rows:
            raise SheetTooLargeError(
                f"Sheet '{ws.title}' exceeds the maximum of {max_rows} rows.")
        rows.append(tuple(row))
    return rows


def read_sheet(path: str, pick: Callable[[Any], Any],
               not_found: Union[str, BaseException]) -> Tuple[List[Tuple], str]:
    """Open the workbook, let ``pick(wb)`` choose the worksheet, materialise
    it and close the workbook: ``(rows, sheet title)``. When ``pick`` returns
    None the workbook is closed and ``not_found`` is raised — a string
    becomes UnrecognizedFormat, an exception instance is raised as is (the
    template reports its missing sheet as a TemplateError)."""
    from bom.parsers import UnrecognizedFormat     # package imports us
    wb = load_workbook_safe(path)
    try:
        ws = pick(wb)
        if ws is None:
            raise not_found if isinstance(not_found, BaseException) else UnrecognizedFormat(not_found)
        return sheet_matrix(ws), ws.title
    finally:
        wb.close()


def head_rows(ws, n: int = 16, cols: int = 24) -> List[Tuple]:
    """First n rows only — what detection needs, without reading the sheet."""
    out = []
    for row in ws.iter_rows(min_row=1, max_row=n, max_col=cols, values_only=True):
        out.append(tuple(row))
    return out


def cell(rows: Sequence[Tuple], r: int, c: int) -> str:
    """1-based (row, col) -> s(value); out of range -> ''."""
    if r < 1 or r > len(rows):
        return ''
    row = rows[r - 1]
    if c < 1 or c > len(row):
        return ''
    return s(row[c - 1])


def raw(rows: Sequence[Tuple], r: int, c: int) -> Any:
    if r < 1 or r > len(rows):
        return None
    row = rows[r - 1]
    if c < 1 or c > len(row):
        return None
    return row[c - 1]


def raw_text_from_rows(rows: Sequence[Tuple], limit: int = 60000) -> str:
    """Tab-joined dump of the non-blank rows, capped — the rawText field is
    for humans reviewing a flagged check, not for the validator."""
    lines = []
    size = 0
    for row in rows:
        vals = [s(v) for v in row]
        if not any(vals):
            continue
        line = '\t'.join(vals).rstrip('\t')
        size += len(line) + 1
        if size > limit:
            lines.append('…')
            break
        lines.append(line)
    return '\n'.join(lines)


# ─── Dell 30-character wrap artefact ─────────────────────────────────────────

_WRAP = 30


def dewrap_dell(text: str) -> str:
    """Undo the service-tag export's hard wrap.

    The exporter cuts the option text into 30-character chunks, right-trims
    each chunk and joins them with a single space. Verified against the three
    archived exports: 'UEFI BIOS Boot Mode with GPT P artition', 'C13 , 3M',
    '10GbE  BASE-T' (original had a space at the chunk start) and '512e 3.5in'
    (original chunk ended in a space, so the join space *is* the original
    space) all reproduce from exactly that model. Walking 30 characters at a
    time therefore recovers the original: after a full chunk that does not
    end in a space, the next character is the inserted join space and is
    removed; a chunk that ends in a space was trimmed-then-rejoined, so the
    following character already belongs to the next chunk.

    Only apply this to the Dell service-tag shape — on unwrapped text it
    would eat a legitimate space every 30 characters.
    """
    out = []
    pos = 0
    n = len(text)
    while pos < n:
        chunk = text[pos:pos + _WRAP]
        out.append(chunk)
        pos += _WRAP
        if len(chunk) == _WRAP and not chunk.endswith(' ') and pos < n and text[pos] == ' ':
            pos += 1
    return ''.join(out)


# ─── server model ─────────────────────────────────────────────────────────────

_MODEL_PATTERNS = [
    (re.compile(r'\bThinkSystem\s+(S[RT]\d{3}[a-z]?)\s*(V\d)\b', re.I),
     lambda m: 'ThinkSystem %s %s' % (m.group(1).upper(), m.group(2).upper())),
    (re.compile(r'\b(S[RT]\d{3})\s*(V\d)\b'),
     lambda m: 'ThinkSystem %s %s' % (m.group(1), m.group(2))),
    (re.compile(r'\b(PowerEdge\s+[A-Z]{1,2}\d{3,4}[A-Za-z]*)\b'),
     lambda m: m.group(1)),
    (re.compile(r'\bDell\s+(R\d{3}[A-Za-z]*)\b'),
     lambda m: 'PowerEdge ' + m.group(1)),
    (re.compile(r'\b(ProLiant\s+[A-Z]{2}\d{3,4}[a-z]*(?:\s+Gen\s?\d+)?)', re.I),
     lambda m: m.group(1)),
    (re.compile(r'\b(SYS-[0-9A-Z]+(?:-[0-9A-Z]+)*)\b'),
     lambda m: m.group(1)),
]


def server_model_from_text(text: Optional[str]) -> Optional[str]:
    """'Server : ThinkSystem SR650 V4-3yr Base Warranty' -> 'ThinkSystem SR650 V4';
    'PowerEdge R760 Server' -> 'PowerEdge R760'; 'Dell R750XS 12x 3.5in LFF'
    -> 'PowerEdge R750XS'; None when nothing recognisable is present."""
    if not text:
        return None
    for rx, fmt in _MODEL_PATTERNS:
        m = rx.search(text)
        if m:
            return fmt(m)
    return None


_SERVER_SUFFIX = re.compile(r'\s+Server\s*$', re.I)


def strip_server_suffix(text: str) -> Optional[str]:
    """'PowerEdge R750 Server' -> 'PowerEdge R750' (the Dell 210- base line
    as the exports write it); None when nothing is left."""
    return _SERVER_SUFFIX.sub('', text or '').strip() or None


def model_from_server_line(desc: str) -> Optional[str]:
    """Server model from a Dell 210- base line: the recognised model when
    there is one, else the line minus its ' Server' suffix."""
    return server_model_from_text(desc) or strip_server_suffix(desc)


# ─── junk / absence / category rules ─────────────────────────────────────────

# Absence indicators the rules engine relies on (rules.is_absence_indicator,
# is_boss_card, is_controller). These win over DROP: a BOM that explicitly
# says 'No BOSS Card' must keep saying so.
#
# German ('Ohne Festplatte', 'Kein Controller', 'Keine BOSS-Karte') and the
# 'Platzhalter' / 'Leermodul' blanks arrived with the Dell solution exports
# (parsers/dell_solution.py), the same way the Portuguese did with Lenovo:
# read as hardware, 'Ohne Festplatte' becomes a diskless node's phantom disk
# and 'Keine BOSS-Karte' becomes a BOSS card nobody quoted.
ABSENCE = re.compile(
    r'^(?:No\s+(?:BOSS|Controller|HBA|RAID|PERC|OCP|Additional Processor|Second Processor|'
    r'Hard Drive|Rear Storage|GPU|Internal|Trusted Platform|TPM)\b'
    r'|(?:Assembly\s+)?BOSS Blank|LOM Blank|Riser Blank|OCP Blank|.*\bBlank\s*$'
    r'|Unconfigured RAID|C\d+,\s*No RAID|No RAID|Select Storage devices'
    r'|Dispositivos de armazenamento'
    r'|Ohne\s+(?:BOSS|Festplatte|Controller|HBA|RAID|PERC|OCP|Betriebssystem|GPU|'
    r'Laufwerk|zweiten|weiteren)'
    r'|Kein(?:e|en|er)?\s+(?:BOSS|Festplatte|Controller|HBA|RAID|PERC|OCP|'
    r'Betriebssystem|GPU|Laufwerk|Datentr[äa]ger|Bereitstellung|DPUs?)'
    r'|C\d+,?\s*[-–—]?\s*[Kk]ein\s+RAID'
    r'|.*\bPlatzhalter\s*$|.*\bLeermodul\b)',
    re.I)

# Pure order-entry noise. Anchored alternatives first, then whole-word hits
# anywhere, then spec-only memory lines. Support/service SKU families are
# handled by part-number prefix in should_drop().
DROP = re.compile(
    r'(?:^(?:UEFI|Performance BIOS|Performance Optimized|RAID\s+\d|Configuration Services|'
    r'System Box|Custom Configuration|Quick Sync|Power Saving|Dell Connectivity|'
    r'Secure Onboarding|Secured Component|Server Secured|OpenManage|DHCP|'
    r'Additional Processor Selected|No HBM|None Required|Decline Selection|'
    r'Basic Next Business Day|ProSupport|Keep Your Drive|Diversified Supply|'
    r'Data Center Environment|Operating mode selection|Disable IPMI|Trigger MFG|'
    r'Registration only|Top Choice|ERP LOT9|Low voltage|High voltage|N\+N Redundancy|'
    r'Feature Enable TPM|Notice for Advanced Format|Configuration ID|'
    r'Configuration Instruction|Controller 0\d|Months|KYD|Premier|Next Business Day|'
    r'A/C Power Recovery|Power Management|Apenas registro|Redund[âa]ncia|Baixa voltagem|'
    r'Alta voltagem|Aviso para|Desativar IPMI|Ambiente de data center|Computa[çc][ãa]o geral|'
    r'Recurso Habilitar|CPK para|Indicador de servi[çc]o|HV \d+U)\b'
    r'|\b(?:Shipping|Marking|Label|Etiqueta|Luggage|Password|Blanks? for|Documentation|'
    r'Media Required|Operating System|Utility Partition|Placement|Filler|Preenchimento|'
    r'Dummy|Sponge|Esponja|Air Duct|Duto de ar|PKG BOM|PACOTE BOM|Packag(?:e|ing)|Welcome Kit|'
    r'EULA|Preload|Laser service|Custom MFG|Customer Solution Center|Service Label|'
    r'Field Deployment|Deployment Services|Energy Star|Support Services?|Warranty|'
    r'XClarity|XCC|FOD|Root of Trust|Power Efficiency|efici[êe]ncia de energia|LPK|'
    r'iDRAC[ ,]+(?:Legacy|Factory|Group)|Group Manager|Factory Generated|Mylar)\b'
    r'|\bATE\s*$'
    r'|^\d+\s*MT/s\s+RDIMMs?\s*$)',
    re.I)

# Dell SKU families that never carry hardware: services/support (709/865/883/
# 892/989/900), iDRAC-and-flavour info lines (379-) and BIOS boot mode (800-).
_DROP_SKU_PREFIX = re.compile(r'^(?:709|865|883|892|989|900|379|800)-')


def is_absence(description: str) -> bool:
    return bool(ABSENCE.match(description or ''))


def should_drop(description: str, part_number: Optional[str] = None) -> bool:
    """True for rows that carry no hardware meaning. Absence indicators are
    protected first; then the description junk list; then Dell SKU families."""
    desc = (description or '').strip()
    if not desc:
        return True
    if is_absence(desc):
        return False
    if DROP.search(desc):
        return True
    if part_number and _DROP_SKU_PREFIX.match(part_number.strip()):
        return True
    return False


# 'No X' / 'X Blank' lines classify as 'other' regardless of keywords so
# 'No BOSS Card' never counts as a BOSS card.
NEGATED = re.compile(
    r'^(?:No\b|None\b|Nenhum\b|Ohne\b|Kein(?:e|en|er)?\b|Assembly BOSS Blank|BOSS Blank|'
    r'LOM Blank|Riser Blank|.*\bBlank\s*$|Decline|C\d+,\s*No RAID|Unconfigured|'
    r'no configured|Select Storage devices|Dispositivos de armazenamento|'
    r'.*\bPlatzhalter\s*$|.*\bLeermodul\b|Auswahl ablehnen)',
    re.I)

# A line that IS the server (the chassis/base line) — checked before the
# 'other' rule so 'Dell R750XS 12x 3.5in LFF, Riser 4 Config' is not 'other'
# because of 'Riser'.
SERVER_LINE = re.compile(
    r'^(?:Dell\s+)?PowerEdge\s+[A-Z]{1,2}\d{3,4}[A-Za-z]*(?:\s+Server)?\s*$'
    r'|^Dell\s+R\d{3}[A-Za-z]*\b.*\b(?:Configure to Order|LFF|SFF|Server)\b'
    r'|^ThinkSystem\s+S[RT]\d{3}\s*V\d\s*$'
    r'|^(?:HPE\s+)?ProLiant\s+[A-Z]{2}\d{3}',
    re.I)

_CAP = r'\d+(?:[.,]\d+)?\s*(?:TB|GB)\b'
_DRV = (r'(?:SSD|HDD|NVMe|NVME|Hard Drive|Hard Disk|Solid State|Drive|Disk|'
        r'disco r[íi]gido|estado s[óo]lido)')

RULES = [
    ('other', re.compile(
        r'\b(?:Label|Filler|Dummy|Blank|Sponge|Air Duct|Cable|CBL|Cabo|Cord|Clip|Clipe|Latch|Trava|'
        r'Luggage|Bezel|Rails?|ReadyRails|Trilho|Heatsink|Heat Sink|Dissipador|Fans?|Ventilador|'
        r'Power Supply|Fonte de alimenta[çc][ãa]o|PSU|Riser|Cage|Compartimento|TPM|'
        r'Trusted Platform|iDRAC|XClarity|XCC|Shipping|Marking|Warranty|Support|Services?|'
        r'License|Licen[çc]a|Subscription|Placement|Mechanical Parts|Bracket|Tray|Foam|'
        r'Interposer|Duct|Duto|Placa do processador|Placa de E/S|MB|'
        r'Board\b(?!.*\bLOM\b)|Motherboard(?!.*\bLOM\b))\b', re.I)),
    ('boss', re.compile(
        r'\b(?:BOSS(?:-[NS]\d)?|M\.2 RAID|B540p|M\.2 Mirroring|M\.2 .*Adapter|'
        r'Boot Optimized)\b', re.I)),
    ('gpu', re.compile(
        r'\b(?:GPU|NVIDIA|Tesla|RTX|L40S?|L4|A\d{2}|H100|H200|Instinct)\b', re.I)),
    ('controller', re.compile(
        # Dell writes a PERC both as 'PERC H755 SAS Front' and bare, as
        # 'H965i Adapter Low Profile' (Ingram Micro item list, 2026-09-30).
        r'\b(?:HBA\s?\d{3}[a-z]*|PERC\s?H\d{3}[A-Z]?|H[3-9]\d{2}[a-zA-Z]?\s+(?:Adapter|Front|Controller|SAS)|'
        r'4[34]0-\d+[ie]|4350-\d+[ie]|[59][34]0-\d+[ie]|'   # Lenovo RAID 530/540/930/940-8i
        r'9\d{3}-\d+[ie]|SAS3?\s+HBA|HBA ThinkSystem|RAID\s+(?:Controller|Adapter|Card)|'
        r'Storage Controller|MR216i|MR416i|AOC-S38\d\d|AOC-S3008|HBA\s*$)\b', re.I)),
    ('nic', re.compile(
        r'\b(?:Ethernet Adapter|Adaptador Ethernet|Ethernet|Network (?:Card|Adapter|Interface)|'
        r'OCP NIC|OCP Ethernet|rNDC|LOM|\d+\s*GbE|\d+\s*GBase-T|\d+\s*GB\s*Base-T|SFP28|SFP\+|'
        r'X710|E810|E610|XL710|X557|I350|5750\d|5741[46]|5720|ConnectX|NIC|BCM5\d{4})\b', re.I)),
    ('chassis', re.compile(r'\b(?:Chassis|Chassi|Backplane|Media Bay)\b', re.I)),
    ('cpu', re.compile(r'\b(?:Xeon|EPYC|Processor|Processador|Ryzen)\b', re.I)),
    ('memory', re.compile(r'\b(?:RDIMM|UDIMM|LRDIMM|TruDDR\d|DDR[45]|DIMM|MRDIMM)\b', re.I)),
    ('storage', re.compile(r'(?:' + _CAP + r'.*\b' + _DRV + r'\b)|(?:\b' + _DRV + r'\b.*' + _CAP + r')')),
    ('chassis', re.compile(
        r'(?:\bServer\s*$|^(?:Dell\s+)?PowerEdge\s+\w+|^ThinkSystem\s+S[RT]\d{3}|\bBase Server\b|'
        r'\bSYS-\d{3}|Optimized System|X13SCH|^(?:HPE\s+)?ProLiant|^Dell\s+R\d{3})', re.I)),
]


def categorize(description: str) -> str:
    """Nine-way classification of a line item description (survey §4).
    First hit wins; order matters and is documented in RULES."""
    d = (description or '').strip()
    if not d:
        return 'other'
    if NEGATED.match(d):
        return 'other'
    if SERVER_LINE.match(d):
        return 'chassis'
    for cat, rx in RULES:
        if rx.search(d):
            return cat
    return 'other'


def is_m2_media(description: str) -> bool:
    """M.2 boot sticks are not data drives: SC//Design's validator ignored
    them (the Arrow expected findings carry no 2-drive/flash-ratio noise)."""
    return re.search(r'\bM\.2\b', description or '') is not None


def demote_m2(category: str, description: str) -> str:
    """'storage' that is M.2 media becomes 'other' — in every parser, whether
    the category came from the description or from a module name."""
    if category == 'storage' and is_m2_media(description):
        return 'other'
    return category


def categorize_line(description: str) -> str:
    """categorize() with the M.2 demotion applied: what a parser calls on a
    line item. categorize() itself stays the raw nine-way rule."""
    return demote_m2(categorize(description), description)


# ─── component accumulator ───────────────────────────────────────────────────

# How two rows of one config are recognised as the same part:
#   'part_or_desc'  part number, else the description (SKU and feature-code
#                   exports: the same SKU on two rows is one part, rows
#                   without one compare by their text)
#   'part+desc'     part number AND description together (hand-typed and
#                   distributor lists, where one reseller number can carry
#                   two different lines)
_ACC_KEYS = ('part_or_desc', 'part+desc')


class ComponentAcc:
    """Per-config component accumulator: sums duplicate rows, keeps first-seen
    order. ``multiplier`` turns per-node quantities into config totals (the
    Dell solution export quotes 8 DIMMs *per server*)."""

    def __init__(self, key: str = 'part_or_desc', multiplier: int = 1):
        if key not in _ACC_KEYS:
            raise ValueError('unknown accumulator key %r' % key)
        self.key = key
        self.multiplier = multiplier
        self.order: List[str] = []
        self.items: Dict[str, BOMComponent] = {}

    def _key(self, part: Optional[str], desc: str) -> str:
        if self.key == 'part_or_desc':
            return part or ('desc:' + desc)
        return (part or '') + '|' + desc

    def add(self, part: Optional[str], desc: str, qty: int, category: str) -> BOMComponent:
        key = self._key(part, desc)
        total = qty * self.multiplier
        comp = self.items.get(key)
        if comp is not None:
            comp.quantity += total
        else:
            comp = self.items[key] = BOMComponent(part_number=part or None, description=desc,
                                                  quantity=total, category=category)
            self.order.append(key)
        return comp

    def __bool__(self) -> bool:
        return bool(self.order)

    def components(self) -> List[BOMComponent]:
        return [self.items[k] for k in self.order]


def split_nx(description: str) -> Tuple[int, str]:
    """'8x 16GB RDIMM…' -> (8, '16GB RDIMM…'); no prefix -> (1, description)."""
    m = NX_PREFIX.match((description or '').strip())
    if m:
        return int(m.group(1)), m.group(2).strip()
    return 1, (description or '').strip()
