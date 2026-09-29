from openpyxl import load_workbook
from xlsx_utils import (sheet_rows as _sheet_rows, to_float as _float,
                        to_int as _int)
from parser_common import build_summary as _build_summary
from cluster_split import cluster_summaries as _cluster_summaries


# Nutanix Collector (vCenter mode) export. It borrows RVTools' sheet names
# (vInfo, vCPU, vMemory, vDisk, vPartition, ...) but not its columns, so it must
# be detected before the RVTools check — a file with "vInfo" otherwise parses
# as RVTools and every VM reads back as 0 vCPU / 0 RAM / powered off (issue #32).
# The flat "vmList" sheet carries the per-VM configuration; utilisation is the
# collector's 95th percentile over its performance window.

def is_nutanix_collector(sheetnames):
    return "vmList" in sheetnames and "vHosts" in sheetnames


def parse_nutanix_collector(file_path):
    wb = load_workbook(file_path, read_only=True, data_only=True)
    hosts = _parse_hosts(wb)
    vms, vm_iops = _parse_vms(wb)
    result = {
        "project": _parse_metadata(wb),
        "hosts": hosts,
        "host_performance": _parse_host_perf(wb, vm_iops),
        "datastores": _parse_datastores(wb),
        "vms": vms,
        "host_nics": _parse_nics(wb),
    }
    wb.close()

    result["summary"] = _build_summary(result, source="nutanix")
    result["clusters"] = _cluster_summaries(result, _build_summary)
    return result


def _parse_metadata(wb):
    rows = list(_sheet_rows(wb, "Metadata"))
    if not rows:
        return {}
    return {k: v for k, v in rows[0].items() if k and v not in (None, "")}


def _parse_hosts(wb):
    hosts = []
    for r in _sheet_rows(wb, "vHosts"):
        speed_mhz = _float(r.get("CPU Speed", 0))
        sockets = _int(r.get("CPUs", 0))
        total_cores = _int(r.get("CPU Cores", 0)) or sockets * _int(r.get("Cores per CPU", 0))
        mem_gib = _float(r.get("Memory Size", 0))

        hosts.append({
            "name": r.get("Host Name", ""),
            "cluster": r.get("Cluster Name", ""),
            "manufacturer": r.get("Vendor", ""),
            "model": r.get("Model", ""),
            "cpu_sockets": sockets,
            "cpu_cores": total_cores,
            # No hyperthreading column in the export; count cores only.
            "cpu_threads": total_cores,
            "cpu_desc": r.get("CPU Model", ""),
            "cpu_ghz": round(speed_mhz / 1000, 3),
            "net_ghz": round(speed_mhz / 1000 * total_cores, 1),
            "memory_kib": mem_gib * 1024 * 1024,
            "memory_gb": round(mem_gib, 1),
            "local_capacity_gib": 0,
            "vm_count": _int(r.get("VMs", 0)),
            "nic_count": _int(r.get("NICs", 0)),
        })
    return hosts


def _parse_host_perf(wb, vm_iops):
    """Host CPU/memory usage from vHosts. IOPS come from the per-VM 95th
    percentile summed per host: the host-level IOPS column only covers the
    host's own devices, while the VM sum reconciles with the vCluster total."""
    perfs = []
    for r in _sheet_rows(wb, "vHosts"):
        name = r.get("Host Name", "")
        cpu_pct = _float(r.get("CPU Usage", 0))
        mem_pct = _float(r.get("Memory Usage", 0))
        total_cores = _int(r.get("CPU Cores", 0))
        total_ghz = _float(r.get("CPU Speed", 0)) / 1000 * total_cores
        mem_mib = _float(r.get("Memory Size", 0)) * 1024
        cpu_ghz = round(total_ghz * cpu_pct / 100, 1)
        mem_used_mib = round(mem_mib * mem_pct / 100, 1)

        perfs.append({
            "host": name,
            "peak_cpu_pct": cpu_pct,
            "peak_cpu_ghz": cpu_ghz,
            "avg_cpu_pct": cpu_pct,
            "avg_cpu_ghz": cpu_ghz,
            "peak_mem_pct": mem_pct,
            "peak_mem_mib": mem_used_mib,
            "avg_mem_pct": mem_pct,
            "avg_mem_mib": mem_used_mib,
            # Only a 95th percentile is collected; peak/avg stay unknown (0)
            # rather than being guessed from it.
            "peak_iops": 0,
            "avg_iops": 0,
            "p95_iops": round(vm_iops.get(name, 0), 1),
            "peak_throughput_mbs": 0,
            "avg_throughput_mbs": 0,
        })
    return perfs


def _parse_datastores(wb):
    """Accessible datastores. One mounted on a single host is local storage
    (boot/scratch VMFS); anything shared is cluster storage."""
    stores = []
    for r in _sheet_rows(wb, "Datastore"):
        if str(r.get("Accessible", "")).lower() != "true":
            continue
        stores.append({
            "name": r.get("Datastore Name", ""),
            "type": "local" if _int(r.get("Host Count", 0)) <= 1 else "cluster",
            "capacity_gib": round(_float(r.get("Capacity (MiB)", 0)) / 1024, 1),
            "used_gib": round(_float(r.get("Consumed (MiB)", 0)) / 1024, 1),
            "free_gib": round(_float(r.get("Freespace (MiB)", 0)) / 1024, 1),
            "vm_count": _int(r.get("VM Count", 0)),
            "mounts": [h for h in str(r.get("Host Names") or "").split(",") if h],
        })
    return stores


def _parse_vms(wb):
    """Returns (vms, p95 IOPS per host of the powered-on VMs)."""
    mem_p95 = {}
    for r in _sheet_rows(wb, "vMemory"):
        mem_p95[r.get("VM Name", "")] = _float(r.get("95th Percentile % (recommended)", 0))

    iops = {}
    for r in _sheet_rows(wb, "vInfo"):
        iops[r.get("VM Name", "")] = _float(r.get("95th Percentile IOPS", 0))

    guest = {}
    for r in _sheet_rows(wb, "vPartition"):
        agg = guest.setdefault(r.get("VM Name", ""), [0.0, 0.0])
        agg[0] += _float(r.get("Capacity (MiB)", 0))
        agg[1] += _float(r.get("Consumed (MiB)", 0))

    vms = []
    host_iops = {}
    for r in _sheet_rows(wb, "vmList"):
        name = r.get("VM Name", "")
        powered_on = str(r.get("Power State", "")).lower() == "poweredon"
        is_template = str(r.get("Template", "")).lower() == "true"
        mem_mib = _float(r.get("Memory (MiB)", 0))
        cap_mib = _float(r.get("Capacity (MiB)", 0))
        guest_cap, guest_used = guest.get(name, (0.0, 0.0))

        # Datastore consumption is blank for some VMs (no tools / no stats).
        # A thick disk consumes its full size; for a thin one fall back to what
        # the guest reports, else the provisioned size (never under-size).
        consumed = r.get("Consumed (MiB)")
        if consumed not in (None, ""):
            used_mib = _float(consumed)
        elif str(r.get("Thin Provisioned", "")).lower() != "true":
            used_mib = cap_mib
        else:
            used_mib = guest_used or cap_mib

        used_mem_gb = round(mem_mib * mem_p95.get(name, 0) / 100 / 1024, 2)
        host = r.get("Host", "")
        if powered_on and not is_template:
            host_iops[host] = host_iops.get(host, 0) + iops.get(name, 0)

        vms.append({
            "name": name,
            "powered_on": powered_on,
            "is_template": is_template,
            "os": r.get("Guest OS") or "",
            "vcpus": _int(r.get("vCPUs", 0)),
            "provisioned_memory_gb": round(mem_mib / 1024, 2),
            "used_memory_gb": used_mem_gb,
            "consumed_memory_gb": used_mem_gb,
            "disk_capacity_gb": round((guest_cap or cap_mib) / 1024, 2),
            "disk_used_gb": round((guest_used or used_mib) / 1024, 2),
            "vdisk_size_gb": round(cap_mib / 1024, 2),
            "vdisk_used_gb": round(used_mib / 1024, 2),
            "datastore": r.get("Datastore", "") or "",
            "host": host,
            "cluster": r.get("Cluster Name", ""),
        })
    return vms, host_iops


def _parse_nics(wb):
    nics = []
    for r in _sheet_rows(wb, "vNICs"):
        nics.append({
            "host": r.get("Host", ""),
            "name": r.get("Network Device", ""),
            "speed_mbps": _float(r.get("Speed (Mbps)", 0)),
            "vendor": "",
            "device": r.get("Adapter", "") or r.get("Drivers", ""),
        })
    return nics
