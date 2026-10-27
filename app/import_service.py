"""The assessment-import pipeline, out of the route handler.

What ``/api/import-liveoptics`` did inline (detect the file type, parse, reject
an unusable export, stamp provenance, collect the data-quality warnings) now
lives here so the MCP upload page runs the identical pipeline
(docs/mcp/01-tool-design.md §3.10, §7). Also here, because the MCP server has
no browser to do them in:

  * ``adjust_summary``: the summary after VM exclusions, the Python port of
    the browser's ``computeAdjustedImportSummary`` (static/js/app.js). The
    port covers the whole function, edits and added VMs included, so the
    server and the UI agree on any payload the UI later re-saves; MCP itself
    exposes exclusions only (§2.7).
  * ``suggested_exclusion``: the UI's patterns for VMs that never migrate
    (the old HCI's storage controller VMs, vSphere's own agent and vCenter
    VMs), so the MCP VM list can flag them the way the exclusion modal does.
  * ``cluster_dataset``: one source cluster of a multi-cluster export, cut
    the way ``cluster_split`` cuts it for the per-cluster summaries.

Not in ``fingerprint.ENGINE_MODULES`` or ``PARSER_MODULES``: the parsers and
``parser_common.build_summary`` are called, never changed, and a sizing is
fingerprinted on its summary, not on the route that produced it.
"""
import copy
import math
import re
from datetime import datetime, timezone

from cluster_split import split_by_cluster
from fingerprint import PARSER_VERSION
from import_checks import build_import_warnings, import_rejection
from liveoptics import parse_liveoptics
from nutanix_collector import is_nutanix_collector, parse_nutanix_collector
from parser_common import collection_date
from request_utils import sha256_file
from rvtools import parse_rvtools

UNRECOGNISED_MESSAGE = ("Unrecognised file format. Please upload an environment "
                        "assessment export (Live Optics, RVTools or Nutanix Collector).")
PARSE_FAILED_MESSAGE = ("Could not parse the file. Upload a valid environment "
                        "assessment export (.xlsx).")


class AssessmentError(Exception):
    """A file the pipeline will not size. ``status`` is the HTTP status the
    REST route answers with, ``code``/``params`` the translatable reason
    (``import_checks.import_rejection``'s shape) when there is one."""

    def __init__(self, message, status=400, code=None, params=None):
        super().__init__(message)
        self.status = status
        self.code = code
        self.params = params or {}


def detect_file_type(path):
    """``rvtools``, ``liveoptics``, ``nutanix`` or None, from the sheet names."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    sheets = set(wb.sheetnames)
    wb.close()
    # Nutanix Collector reuses RVTools sheet names (vInfo, ...) — check it first.
    if is_nutanix_collector(sheets):
        return "nutanix"
    if "vInfo" in sheets or "vMetaData" in sheets:
        return "rvtools"
    if "ESX Hosts" in sheets or "Details" in sheets:
        return "liveoptics"
    return None


PARSERS = {
    "rvtools": parse_rvtools,
    "liveoptics": parse_liveoptics,
    "nutanix": parse_nutanix_collector,
}


def _utcnow_iso():
    return datetime.now(timezone.utc).isoformat()


def source_meta_for(data, file_name, file_type, file_digest):
    """Provenance for the project view and the export appendix
    (docs/projects-plan.md §8): which file, when, read by which parser. The
    digest lets a client warn when the same file is imported into a project
    twice."""
    return {
        "file_name": file_name,
        "file_type": file_type,
        "file_sha256": file_digest,
        "imported_at": _utcnow_iso(),
        # When the assessment was collected, from the file itself; a
        # multi-file import keeps the newest copy of an overlap.
        "collected_at": collection_date(data.get("project")),
        "host_count": len(data.get("hosts") or []),
        "vm_count": len(data.get("vms") or []),
        "parser_version": PARSER_VERSION,
    }


def parse_assessment(path, file_name):
    """One export file to a dataset the route, the merge and the MCP upload
    all consume: the parser's records plus ``file_type``, ``source_meta`` and
    ``import_warnings``. Raises AssessmentError for a file that is not an
    assessment (400) or an export too incomplete to size (422); any other
    failure is the parser's and propagates for the caller to log."""
    file_digest = sha256_file(path)
    file_type = detect_file_type(path)
    parser = PARSERS.get(file_type)
    if parser is None:
        raise AssessmentError(UNRECOGNISED_MESSAGE)
    data = parser(path)

    # Too incomplete to size at all (e.g. an aborted Live Optics run).
    # The code lets a client show a translated message.
    rejection = import_rejection(data, file_type)
    if rejection:
        raise AssessmentError(
            "This Live Optics export is incomplete: its datastores account for only "
            f"{rejection['params']['covered']} of {rejection['params']['vms']} VMs. "
            "The collection was probably stopped before its first complete snapshot. "
            "Re-run Live Optics and let it complete before exporting.",
            status=422, code=rejection["code"], params=rejection["params"])

    data["file_type"] = file_type
    data["source_meta"] = source_meta_for(data, file_name, file_type, file_digest)
    data["import_warnings"] = build_import_warnings(data, file_type)
    return data


def cluster_dataset(data, name):
    """The records of one source cluster of a multi-cluster dataset, in the
    dataset shape (hosts, vms, perf, datastores, NICs), or None when the
    dataset holds no such cluster. Single-cluster datasets have no split: the
    aggregate is the cluster."""
    split = split_by_cluster(data) or []
    for cluster, sub in split:
        if cluster == name:
            return sub
    return None


# ── VM exclusions (the browser's computeAdjustedImportSummary) ───────────────

def _js_round(value, digits):
    """``Math.round(x * 10**d) / 10**d``: half away from zero for positives
    (Python's round() is half-to-even, so 2.25 -> 2.2 there and 2.3 here)."""
    factor = 10 ** digits
    return math.floor(value * factor + 0.5) / factor


def _num(value):
    try:
        v = float(value)
    except (TypeError, ValueError):
        return 0.0
    return 0.0 if math.isnan(v) else v


def _is_active(vm):
    return bool(vm.get("powered_on")) and not vm.get("is_template")


def adjust_summary(base, vms, exclude_compute=(), exclude_storage=(),
                   include_local_storage=False, vm_config=None,
                   added=(), removed=()):
    """``base`` after leaving VMs out, exactly as the browser computes it
    (``computeAdjustedImportSummary``): compute and RAM totals, the active
    count, the largest VM and the current ratio are rebuilt from the included
    powered-on VMs; storage-excluded VMs' disk is subtracted from the measured
    datastore totals; local storage is folded in on request.

    Indices are positions in ``vms``. ``vm_config`` maps an index to per-VM
    edits (vcpus, provisioned_memory_gb, vdisk_used_gb); ``added`` and
    ``removed`` are the browser's hand-added VMs and their removal, kept so a
    re-saved UI payload adjusts the same way. With no exclusions this
    reproduces the import totals."""
    exclude_compute = {int(i) for i in exclude_compute}
    exclude_storage = {int(i) for i in exclude_storage}
    added = {int(i) for i in added}
    removed = {int(i) for i in removed}
    config = {}
    for key, overrides in (vm_config or {}).items():
        try:
            config[int(key)] = overrides or {}
        except (TypeError, ValueError):
            continue

    def val(vm, i, field):
        ov = config.get(i)
        return ov[field] if ov and field in ov else vm.get(field)

    adjusted = copy.deepcopy(base)

    # Recompute compute/RAM totals straight from the (possibly user-edited)
    # VM list rather than subtracting from the original, so per-VM edits
    # flow through.
    sum_vcpus = sum_prov_ram = sum_used_ram = 0.0
    active_included = 0
    max_vm_ram = max_vm_cores = 0.0
    for i, vm in enumerate(vms):
        if not _is_active(vm) or i in exclude_compute:
            continue
        cores = _num(val(vm, i, "vcpus"))
        ram = _num(val(vm, i, "provisioned_memory_gb"))
        sum_vcpus += cores
        sum_prov_ram += ram
        sum_used_ram += _num(vm.get("consumed_memory_gb"))
        active_included += 1
        max_vm_ram = max(max_vm_ram, ram)
        max_vm_cores = max(max_vm_cores, cores)

    adjusted["total_vcpus"] = _int_if_whole(sum_vcpus)
    adjusted["total_vm_provisioned_memory_gb"] = _js_round(sum_prov_ram, 1)
    adjusted["total_vm_used_memory_gb"] = _js_round(sum_used_ram, 1)
    adjusted["active_vms"] = active_included

    # Server-level scans have no real overcommit (hosts == VMs), so keep the
    # assumed 3:1 default instead of recomputing a meaningless 1:1.
    host_cores = _num(adjusted.get("total_host_cores"))
    if not adjusted.get("vcpu_ratio_assumed") and host_cores > 0:
        adjusted["vcpu_per_core_ratio"] = _js_round(sum_vcpus / host_cores, 2)

    adjusted["max_vm_ram_gb"] = _int_if_whole(max_vm_ram)
    adjusted["max_vm_cores"] = _int_if_whole(max_vm_cores)

    # Excluded/removed IMPORTED VMs are subtracted from the measured datastore
    # totals (added VMs aren't in those totals, so they're skipped here).
    excl_all = excl_prov_active = excl_active = 0.0
    for i in exclude_storage:
        if i in added or i >= len(vms):
            continue
        vm = vms[i]
        excl_all += _num(vm.get("vdisk_used_gb"))
        if _is_active(vm):
            excl_active += _num(vm.get("vdisk_used_gb"))
            excl_prov_active += _num(vm.get("vdisk_size_gb"))

    # Added VMs contribute net-new storage (unless removed or storage-
    # excluded). A new VM has one storage figure, used for both used and
    # provisioned.
    add_all = add_prov_active = add_active = 0.0
    for i in added:
        if i in removed or i in exclude_storage or i >= len(vms):
            continue
        vm = vms[i]
        used = _num(val(vm, i, "vdisk_used_gb"))
        add_all += used
        if _is_active(vm):
            add_active += used
            add_prov_active += used

    adjusted["datastore_used_tb"] = _js_round(
        _num(base.get("datastore_used_tb")) - excl_all / 1024 + add_all / 1024, 2)
    prov_gb = _js_round(
        _num(base.get("total_vm_provisioned_storage_gb")) - excl_prov_active + add_prov_active, 1)
    adjusted["total_vm_provisioned_storage_gb"] = prov_gb
    adjusted["total_vm_provisioned_storage_tb"] = _js_round(prov_gb / 1024, 2)
    used_gb = _js_round(
        _num(base.get("total_vm_used_storage_gb")) - excl_active + add_active, 1)
    adjusted["total_vm_used_storage_gb"] = used_gb
    adjusted["total_vm_used_storage_tb"] = _js_round(used_gb / 1024, 2)

    if adjusted["datastore_used_tb"] < 0:
        adjusted["datastore_used_tb"] = 0

    # Per-host local storage (ISOs/templates etc.) — added to the cluster
    # basis only when the user opts in.
    if include_local_storage:
        adjusted["datastore_used_tb"] = _js_round(
            adjusted["datastore_used_tb"] + _num(base.get("local_used_tb")), 2)
        adjusted["datastore_total_tb"] = _js_round(
            _num(adjusted.get("datastore_total_tb")) + _num(base.get("local_total_tb")), 2)

    return adjusted


def _int_if_whole(value):
    """The parsers give whole numbers as ints; keep them that way after the
    float arithmetic so a summary round-trips through JSON unchanged."""
    return int(value) if float(value).is_integer() else value


# ── suggested exclusions (the exclusion modal's patterns) ────────────────────

# Storage controller VMs of the HCI being replaced: they never move.
CVM_PATTERNS = [re.compile(p, re.I) for p in (
    r"^stctlvm", r"^scvm", r"^cvm", r"^ntnx", r"^nutanix",
    r"^svt-", r"^omnistack", r"simplivity", r"^vsa-", r"^vsa\d",
    r"^hpe\s*simplivity", r"witness",
)]
# VMware's own infrastructure VMs, which never move to the target platform:
# vCLS-<uuid> agent VMs (vSphere Cluster Services, one or more per cluster
# since vSphere 7.0u1) and vCenter Server appliances (vcenter/VCSA naming).
VMWARE_ARTIFACT_PATTERNS = [re.compile(p, re.I) for p in (
    r"^vcls-", r"vcenter", r"^vcsa",
)]


def suggested_exclusion(name):
    """``"cvm"``, ``"vmware_infrastructure"`` or None for a VM name."""
    n = str(name or "")
    if any(p.search(n) for p in CVM_PATTERNS):
        return "cvm"
    if any(p.search(n) for p in VMWARE_ARTIFACT_PATTERNS):
        return "vmware_infrastructure"
    return None
