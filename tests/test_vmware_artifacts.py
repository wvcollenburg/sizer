"""The Remove-VMware-artifacts patterns, tested against a real export.

The patterns live in app.js (VMWARE_ARTIFACT_PATTERNS); this test extracts
them from the source so the JS and the expectation cannot drift, then runs
them over the archived 3-cluster LiveOptics file: every vCLS agent VM must
match, and none of the 84 real workload VMs may.

Run: .venv/bin/python -m pytest tests/test_vmware_artifacts.py -q
"""
import os
import re

import pytest

APP_JS = os.path.join(os.path.dirname(__file__), "..", "app", "static", "js", "app.js")
FIXTURE = os.path.join(os.path.dirname(__file__), "..", "_archive",
                       "LiveOptics_3455045_VMWARE_09_07_2026.xlsx")

pytestmark = pytest.mark.skipif(not os.path.exists(FIXTURE),
                                reason="archived LiveOptics file not present")


def _patterns_from_js():
    src = open(APP_JS, encoding="utf-8").read()
    m = re.search(r"VMWARE_ARTIFACT_PATTERNS\s*=\s*\[(.*?)\]", src, re.S)
    assert m, "VMWARE_ARTIFACT_PATTERNS not found in app.js"
    pats = []
    for body, flags in re.findall(r"/((?:[^/\\]|\\.)+)/(\w*)", m.group(1)):
        pats.append(re.compile(body, re.I if "i" in flags else 0))
    assert pats, "no regexes parsed from VMWARE_ARTIFACT_PATTERNS"
    return pats


def _vm_names():
    from openpyxl import load_workbook
    wb = load_workbook(FIXTURE, read_only=True)
    rows = wb["VMs"].iter_rows(values_only=True)
    hdr = list(next(rows))
    i = hdr.index("VM Name")
    names = [r[i] for r in rows if r[i]]
    wb.close()
    return names


def test_patterns_match_all_vcls_and_no_workload_vms():
    pats = _patterns_from_js()
    names = _vm_names()
    hits = {n for n in names if any(p.search(n) for p in pats)}
    vcls = {n for n in names if n.lower().startswith("vcls-")}
    assert vcls, "fixture should contain vCLS agent VMs"
    assert vcls <= hits, f"missed vCLS VMs: {vcls - hits}"
    # Nothing that isn't VMware infrastructure may match.
    false_pos = {n for n in hits
                 if not (n.lower().startswith(("vcls-", "vcsa"))
                         or "vcenter" in n.lower())}
    assert not false_pos, f"false positives: {false_pos}"


def test_patterns_catch_vcenter_naming():
    pats = _patterns_from_js()
    for name in ("vCenter", "VCSA01", "prod-vcenter-01", "vcls-abc"):
        assert any(p.search(name) for p in pats), name
    for name in ("ARX-Rakkestad", "Deploy01", "vc-not-center", "sql-vcuster"):
        assert not any(p.search(name) for p in pats), name
