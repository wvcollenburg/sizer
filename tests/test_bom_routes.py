"""BOM checker HTTP routes (docs/bom-checker-build.md §6).

Covers the contract the UI relies on and the rules that are invisible until
they break: project visibility gates every check, the upload never stores
the file, unrecognised files are refused (not best-effort parsed), a flagged
result opens an admin review the user can see the outcome of, and a re-check
appends history instead of overwriting it.

Run: .venv/bin/python -m pytest tests/test_bom_routes.py -q
"""
import io
import os
import sys

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("ENABLE_SCHEDULER", "0")
os.environ.setdefault("SECRET_KEY", "test-secret")
os.environ.pop("SESSION_COOKIE_SECURE", None)
os.environ.pop("ANTHROPIC_API_KEY", None)
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

import pytest  # noqa: E402
import app as appmod  # noqa: E402
from database import db  # noqa: E402
from extensions import limiter  # noqa: E402
from auth_models import ROLE_SUPER_ADMIN, User  # noqa: E402
from hcl_models import HclComponent, HclPlatform, HclPlatformComponent  # noqa: E402
from bom.normalize import BOMComponent, BOMConfig, NormalizedBOM  # noqa: E402
from bom.hcl_scrape import component_attrs  # noqa: E402

PASSWORD = "Abcdef1!xy"
PARTNER = "pm@partnerco.example"
OTHER = "someone@elsewhere.example"
ADMIN = "admin@scalecomputing.com"


@pytest.fixture()
def app():
    application = appmod.app
    application.config["TESTING"] = True
    application.config["RATELIMIT_ENABLED"] = False
    limiter.enabled = False
    with application.app_context():
        db.drop_all()
        db.create_all()
        _seed_catalog()
    return application


def _comp(kind, part, desc, **attrs):
    a = component_attrs(kind, desc)
    a.update(attrs)
    c = HclComponent(kind=kind, part_number=part, description=desc, attrs=a)
    db.session.add(c)
    return c


def _seed_catalog():
    p = HclPlatform(brand="lenovo", sc_model="HC1450D", server="ThinkSystem SR630V3",
                    form_factor="1U", socket="FCLGA4677", sockets=2)
    db.session.add(p)
    parts = [
        _comp("nic", "4XC7A80269", "ThinkSystem Intel E810-DA4 10/25GbE SFP28 4-Port OCP Ethernet Adapter"),
        _comp("nic", "4XC7A08294", "ThinkSystem Intel E810-DA2 10/25GbE SFP28 2-Port OCP Ethernet Adapter"),
        _comp("hba", "4Y37A78602", "ThinkSystem 440-16i SAS/SATA PCIe Gen4 12Gb HBA"),
        _comp("cpu", "PK8071305120500", "Intel Xeon Gold 6526Y Processor", model="Xeon Gold 6526Y"),
    ]
    db.session.flush()
    for c in parts:
        db.session.add(HclPlatformComponent(platform=p, component=c))
    db.session.commit()


def client_for(app, email):
    c = app.test_client()
    c.post("/api/auth/signup", json={"email": email, "password": PASSWORD, "accept_privacy": True})
    c.post("/api/auth/login", json={"email": email, "password": PASSWORD})
    return c


def promote(app, email):
    with app.app_context():
        u = User.query.filter_by(email=email).first()
        u.role = ROLE_SUPER_ADMIN
        db.session.commit()


def make_project(c, name="Acme"):
    r = c.post("/api/projects/", json={"name": name})
    assert r.status_code == 201, r.get_data(as_text=True)
    return r.get_json()


def _snapshot():
    return {"clusters": [{
        "name": "Prod",
        "summary": {"active_vms": 40, "nic_speed_mbps": 25000, "max_vm_ram_gb": 64,
                    "max_vm_cores": 8, "total_host_ghz": 200},
        "projection": {"years": 5, "base_ram_gb": 300, "base_storage_tb": 10,
                       "compute_floor": {"active": False}},
        "recommendation": {
            "model": "HC1450D", "node_count": 3, "hci_node_count": 3, "cluster_layout": [3],
            "sized_full_cluster": False,
            "totals": {"cores": 90, "ram_gb": 700, "usable_storage_tb": 21.12, "total_ghz": 300},
            "n_minus_1": {"cores": 60, "ram_gb": 466, "usable_storage_tb": 21.12},
            "utilization": {"cpu": {"abs": {"total": 20}}, "ram": {"abs": {"total": 300}},
                            "storage": {"abs": {"total": 12}}},
            "usable_ram_per_node_gb": 233, "threads_per_node": 64, "nic_ports": 4,
            "storage_config": {"drive_counts": {"NVMe": 4}},
        },
        "refs": {"mode": "validated"},
    }], "tunables": "t1"}


def sized(c, project_id, name="Option 1"):
    row = c.post("/api/configs/", json={
        "name": name, "payload": {"mode": "import"}, "project_id": project_id}).get_json()
    r = c.put(f"/api/sizings/{row['id']}/result", json=_snapshot())
    assert r.status_code == 200, r.get_data(as_text=True)
    return row


def _bom(nic="ThinkSystem Intel E810-DA4 10/25GbE SFP28 4-Port OCP Ethernet Adapter"):
    def comp(cat, desc, part, qty):
        return BOMComponent(part_number=part, description=desc, quantity=qty, category=cat)
    return NormalizedBOM(vendor="Lenovo", configs=[BOMConfig(
        name="Acme PROD", server_model="ThinkSystem SR630 V3", node_count=3, components=[
            comp("chassis", 'ThinkSystem V3 1U 10x2.5" Chassis', "BLK4", 3),
            comp("cpu", "Intel Xeon Gold 6526Y 16C 195W 2.8GHz Processor", "BYVX", 3),
            comp("memory", "ThinkSystem 32GB TruDDR5 5600MHz (2Rx8) RDIMM", "BWJC", 24),
            comp("storage", 'ThinkSystem 2.5" U.2 VA 3.84TB Read Intensive NVMe PCIe 4.0 x4 HS SSD', "C18M", 12),
            comp("nic", nic, "BP8L", 3),
        ])])


def template_upload(bom):
    from bom.parsers.template import build_template_bytes
    return (io.BytesIO(build_template_bytes(bom=bom)), "acme-bom.xlsx")


def upload(c, project_id, bom=None, **form):
    data = {"file": template_upload(bom or _bom())}
    data.update(form)
    return c.post(f"/api/projects/{project_id}/bom-checks", data=data,
                  content_type="multipart/form-data")


# ── capabilities / template / pre-fill gating ────────────────────────────────

def test_capabilities_reports_catalog_and_no_prefill_without_a_key(app):
    c = client_for(app, PARTNER)
    d = c.get("/api/bom/capabilities").get_json()
    assert d["ai_prefill_available"] is False
    assert ".xlsx" in d["accepted_extensions"]
    assert d["catalog"]["platforms"] == 1 and d["catalog"]["components"] == 4
    assert d["template_url"] == "/api/bom/template"


def test_template_downloads_as_xlsx(app):
    c = client_for(app, PARTNER)
    r = c.get("/api/bom/template")
    assert r.status_code == 200
    assert r.data[:4] == b"PK\x03\x04"
    assert "sc-bom-template.xlsx" in r.headers["Content-Disposition"]


def test_prefill_is_503_when_not_configured(app):
    c = client_for(app, PARTNER)
    r = c.post("/api/bom/prefill", data={"file": (io.BytesIO(b"hello"), "quote.txt")},
               content_type="multipart/form-data")
    assert r.status_code == 503


def test_routes_require_login(app):
    c = app.test_client()
    assert c.get("/api/bom/capabilities").status_code == 401
    assert c.get("/api/bom/template").status_code == 401


# ── upload → result ──────────────────────────────────────────────────────────

def test_upload_runs_the_check_and_stores_no_file(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    r = upload(c, project["id"], name="Lenovo quote")
    assert r.status_code == 201, r.get_data(as_text=True)
    d = r.get_json()
    assert d["name"] == "Lenovo quote"
    assert d["file_format"] == "template"
    assert d["vendor"] == "Lenovo"
    assert d["technical_verdict"] == "PASS"
    assert d["fit_verdict"] is None
    assert d["review_status"] == "none"
    cr = d["result"]["technical"]["config_results"][0]
    assert cr["platform"]["sc_models"] == ["HC1450D"]
    assert d["normalized"]["configs"][0]["nodeCount"] == 3
    assert d["history"] and d["history"][0]["technical_verdict"] == "PASS"
    listed = c.get(f"/api/projects/{project['id']}/bom-checks").get_json()
    assert [x["id"] for x in listed] == [d["id"]]
    assert "normalized" not in listed[0]


def test_upload_against_a_sizing_adds_the_fit_verdict(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    sizing = sized(c, project["id"])
    r = upload(c, project["id"], sizing_id=str(sizing["id"]))
    assert r.status_code == 201, r.get_data(as_text=True)
    d = r.get_json()
    assert d["configuration_id"] == sizing["id"]
    assert d["sizing_name"] == "Option 1"
    fit = d["result"]["fit"]
    assert fit["sizing"]["id"] == sizing["id"]
    assert fit["verdict"] in ("match", "bigger", "fits", "smaller", "unknown")
    keys = {dim["key"] for dim in fit["dimensions"]}
    assert {"nodes", "cores", "ram", "storage", "nic"} <= keys
    assert d["fit_verdict"] == fit["verdict"]


def test_sizing_must_belong_to_the_project(app):
    c = client_for(app, PARTNER)
    p1 = make_project(c, "A")
    p2 = make_project(c, "B")
    sizing = sized(c, p2["id"])
    r = upload(c, p1["id"], sizing_id=str(sizing["id"]))
    assert r.status_code == 400


def test_unrecognised_file_is_refused_with_a_hint(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["random", "sheet"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = c.post(f"/api/projects/{project['id']}/bom-checks",
               data={"file": (buf, "random.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 400
    assert "hint" in r.get_json()


def test_template_errors_come_back_row_by_row(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    from bom.parsers.template import build_template_bytes
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(build_template_bytes(bom=_bom())))
    ws = wb["BOM"]
    headers = [cell.value for cell in ws[1]]
    cat_col = headers.index("Category") + 1
    ws.cell(row=2, column=cat_col, value="widget")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = c.post(f"/api/projects/{project['id']}/bom-checks",
               data={"file": (buf, "bad.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 400
    d = r.get_json()
    assert d["details"] and any("2" in line for line in d["details"])


def test_non_xlsx_bytes_are_rejected(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    r = c.post(f"/api/projects/{project['id']}/bom-checks",
               data={"file": (io.BytesIO(b"not a workbook"), "x.xlsx")},
               content_type="multipart/form-data")
    assert r.status_code == 400
    r = c.post(f"/api/projects/{project['id']}/bom-checks",
               data={"file": (io.BytesIO(b"%PDF-1.4"), "x.pdf")},
               content_type="multipart/form-data")
    assert r.status_code == 400


# ── visibility ───────────────────────────────────────────────────────────────

def test_other_tenants_see_nothing(app):
    owner = client_for(app, PARTNER)
    project = make_project(owner)
    check = upload(owner, project["id"]).get_json()
    other = client_for(app, OTHER)
    assert other.get(f"/api/projects/{project['id']}/bom-checks").status_code == 404
    assert other.get(f"/api/bom-checks/{check['id']}").status_code == 404
    assert other.delete(f"/api/bom-checks/{check['id']}").status_code == 404
    assert other.post(f"/api/bom-checks/{check['id']}/recheck", json={}).status_code == 404


def test_delete_soft_deletes(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    check = upload(c, project["id"]).get_json()
    assert c.delete(f"/api/bom-checks/{check['id']}").status_code == 200
    assert c.get(f"/api/bom-checks/{check['id']}").status_code == 404
    assert c.get(f"/api/projects/{project['id']}/bom-checks").get_json() == []


# ── review queue ─────────────────────────────────────────────────────────────

def test_unknown_nic_opens_a_review_the_user_can_follow(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    check = upload(c, project["id"], bom=_bom(nic="Mellanox ConnectX-6 Dx 25GbE 2-port OCP")).get_json()
    assert check["technical_verdict"] == "FAIL"
    assert "nic_not_in_hcl" in check["flag_reasons"]
    assert check["review_status"] == "open"
    assert check["result"]["suggestions"][0]["candidates"][0]["part_number"] == "4XC7A08294"

    # partners cannot reach the admin queue
    assert c.get("/admin/api/bom-reviews").status_code == 403

    admin = client_for(app, ADMIN)
    promote(app, ADMIN)
    queue = admin.get("/admin/api/bom-reviews?status=open").get_json()
    assert [q["id"] for q in queue] == [check["id"]]
    assert queue[0]["project"]["name"] == "Acme"
    assert queue[0]["owner_email"] == PARTNER
    r = admin.post(f"/admin/api/bom-reviews/{check['id']}",
                   json={"status": "incorrect", "note": "CX6 Dx is on the HCL as of last week"})
    assert r.status_code == 200
    assert admin.get("/admin/api/bom-reviews?status=open").get_json() == []

    seen = c.get(f"/api/bom-checks/{check['id']}").get_json()
    assert seen["review_status"] == "incorrect"
    assert "CX6" in seen["review_note"]
    assert admin.post(f"/admin/api/bom-reviews/{check['id']}", json={"status": "bogus"}).status_code == 400


# ── re-check ─────────────────────────────────────────────────────────────────

def test_recheck_appends_history_and_can_switch_sizing(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    sizing = sized(c, project["id"])
    check = upload(c, project["id"]).get_json()
    assert check["configuration_id"] is None
    r = c.post(f"/api/bom-checks/{check['id']}/recheck", json={"sizing_id": sizing["id"]})
    assert r.status_code == 200, r.get_data(as_text=True)
    d = r.get_json()
    assert d["configuration_id"] == sizing["id"]
    assert d["result"]["fit"] is not None
    assert len(d["history"]) == 2
    # dropping the sizing again
    d = c.post(f"/api/bom-checks/{check['id']}/recheck", json={"sizing_id": None}).get_json()
    assert d["configuration_id"] is None and d["result"]["fit"] is None
    assert len(d["history"]) == 3
    # without a body the previous sizing is kept
    c.post(f"/api/bom-checks/{check['id']}/recheck", json={"sizing_id": sizing["id"]})
    d = c.post(f"/api/bom-checks/{check['id']}/recheck").get_json()
    assert d["configuration_id"] == sizing["id"]


def test_recheck_drops_a_sizing_moved_out_of_the_project(app):
    # Finding: "recheck reuses a stored sizing after it was moved out of the
    # project" — the implicit branch must enforce the same project-membership
    # rule as the explicit sizing_id path (_sizing_for).
    c = client_for(app, PARTNER)
    p1 = make_project(c, "A")
    p2 = make_project(c, "B")
    sizing = sized(c, p1["id"])
    check = upload(c, p1["id"], sizing_id=str(sizing["id"])).get_json()
    assert check["configuration_id"] == sizing["id"]
    r = c.post(f"/api/sizings/{sizing['id']}/move", json={"project_id": p2["id"]})
    assert r.status_code == 200, r.get_data(as_text=True)
    d = c.post(f"/api/bom-checks/{check['id']}/recheck", json={}).get_json()
    assert d["configuration_id"] is None
    assert d["result"]["fit"] is None
    # ... exactly like the explicit path refuses it.
    r = c.post(f"/api/bom-checks/{check['id']}/recheck", json={"sizing_id": sizing["id"]})
    assert r.status_code == 400


def test_decompression_bomb_upload_gets_a_clear_400(app):
    # Finding: "xlsx decompression bomb: sharedStrings is loaded eagerly
    # before any row cap applies" — the route must refuse it with a clear
    # error, not the generic 'could not process the file' message.
    import zipfile
    c = client_for(app, PARTNER)
    project = make_project(c)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/sharedStrings.xml", b"<si>x</si>" * (6 * 1024 * 1024))  # 60 MB
    buf.seek(0)
    r = c.post(f"/api/projects/{project['id']}/bom-checks",
               data={"file": (buf, "bomb.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 400
    assert "too large" in r.get_json()["error"].lower()


def test_no_price_shaped_keys_in_check_responses(app):
    """Same guard test_security.py applies to /api/models: nothing on the
    wire may look like a price or a tier."""
    from test_security import _offending_keys  # type: ignore
    c = client_for(app, PARTNER)
    project = make_project(c)
    sizing = sized(c, project["id"])
    d = upload(c, project["id"], sizing_id=str(sizing["id"])).get_json()
    assert _offending_keys(d) == []


# ── consent-based retention of rejected files ────────────────────────────────
# The check route never stores an upload. When a file is refused as
# unrecognisable the response carries retainable=true; only the user's second,
# explicit POST lands the bytes, where the super admin can fetch them over
# HTTP to teach the parsers the format. Rows age out after 90 days.

def _random_workbook_bytes():
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["random", "sheet", "nobody", "knows"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_unrecognised_upload_is_marked_retainable_and_stores_nothing(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    r = c.post(f"/api/projects/{project['id']}/bom-checks",
               data={"file": (io.BytesIO(_random_workbook_bytes()), "mystery.xlsx")},
               content_type="multipart/form-data")
    assert r.status_code == 400
    assert r.get_json()["retainable"] is True
    with app.app_context():
        from bom_models import BomRejectedFile
        assert BomRejectedFile.query.count() == 0


def test_template_row_errors_are_not_retainable(app):
    """A fixable template mistake gets row errors, not a retention offer."""
    c = client_for(app, PARTNER)
    project = make_project(c)
    from bom.parsers.template import build_template_bytes
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(build_template_bytes(bom=_bom())))
    ws = wb["BOM"]
    headers = [cell.value for cell in ws[1]]
    ws.cell(row=2, column=headers.index("Category") + 1, value="widget")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = c.post(f"/api/projects/{project['id']}/bom-checks",
               data={"file": (buf, "bad.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 400
    assert "retainable" not in r.get_json()


def test_consent_flow_share_download_delete(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    payload = _random_workbook_bytes()
    r = c.post(f"/api/projects/{project['id']}/bom-rejects",
               data={"file": (io.BytesIO(payload), "mystery.xlsx"),
                     "error": "This file is not a BOM format the checker recognises."},
               content_type="multipart/form-data")
    assert r.status_code == 201, r.get_data(as_text=True)
    row = r.get_json()
    assert row["filename"] == "mystery.xlsx"
    assert row["size_bytes"] == len(payload)
    assert "content" not in row

    # the same file again: one copy is enough
    r2 = c.post(f"/api/projects/{project['id']}/bom-rejects",
                data={"file": (io.BytesIO(payload), "mystery.xlsx")},
                content_type="multipart/form-data")
    assert r2.status_code == 200 and r2.get_json()["id"] == row["id"]

    # partners cannot reach the admin side
    assert c.get("/admin/api/bom-rejects").status_code == 403
    assert c.get(f"/admin/api/bom-rejects/{row['id']}/file").status_code == 403

    admin = client_for(app, ADMIN)
    promote(app, ADMIN)
    listed = admin.get("/admin/api/bom-rejects").get_json()
    assert [x["id"] for x in listed] == [row["id"]]
    assert listed[0]["owner_email"] == PARTNER
    assert listed[0]["project_name"] == "Acme"
    got = admin.get(f"/admin/api/bom-rejects/{row['id']}/file")
    assert got.status_code == 200 and got.data == payload
    assert "mystery.xlsx" in got.headers["Content-Disposition"]

    assert admin.delete(f"/admin/api/bom-rejects/{row['id']}").status_code == 200
    assert admin.get("/admin/api/bom-rejects").get_json() == []
    assert admin.get(f"/admin/api/bom-rejects/{row['id']}/file").status_code == 404


def test_consent_post_respects_project_write_rights(app):
    owner = client_for(app, PARTNER)
    project = make_project(owner)
    outsider = client_for(app, OTHER)
    r = outsider.post(f"/api/projects/{project['id']}/bom-rejects",
                      data={"file": (io.BytesIO(_random_workbook_bytes()), "x.xlsx")},
                      content_type="multipart/form-data")
    assert r.status_code in (403, 404)


def test_consent_post_validates_the_file(app):
    c = client_for(app, PARTNER)
    project = make_project(c)
    r = c.post(f"/api/projects/{project['id']}/bom-rejects",
               data={"file": (io.BytesIO(b"not a zip"), "x.xlsx")},
               content_type="multipart/form-data")
    assert r.status_code == 400
    r = c.post(f"/api/projects/{project['id']}/bom-rejects",
               data={"file": (io.BytesIO(b"%PDF-1.4"), "x.pdf")},
               content_type="multipart/form-data")
    assert r.status_code == 400


def test_rejected_files_age_out_after_retention(app):
    from datetime import timedelta
    c = client_for(app, PARTNER)
    project = make_project(c)
    c.post(f"/api/projects/{project['id']}/bom-rejects",
           data={"file": (io.BytesIO(_random_workbook_bytes()), "old.xlsx")},
           content_type="multipart/form-data")
    with app.app_context():
        from auth import _purge_expired_bom_rejects
        from auth_models import _utcnow
        from bom_models import BomRejectedFile
        row = BomRejectedFile.query.one()
        assert _purge_expired_bom_rejects() == 0          # fresh: kept
        row = BomRejectedFile.query.one()
        row.created_at = _utcnow() - timedelta(days=BomRejectedFile.RETENTION_DAYS + 1)
        db.session.commit()
        assert _purge_expired_bom_rejects() == 1
        db.session.commit()
        assert BomRejectedFile.query.count() == 0
