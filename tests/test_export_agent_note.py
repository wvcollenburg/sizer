"""Exports say so when the quoted hardware came from an agent-read BOM
(owner decision B, 2026-09-29; export_override.resolve → bom_agent_read).

Run: .venv/bin/python -m pytest tests/test_export_agent_note.py -q
"""
import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

from test_export_compute_floor import _compute_driven_rec  # noqa: E402
from i18n import translator  # noqa: E402

NOTE = translator("en")("export.common.agent_bom_note")


def _docx_text(buf):
    import docx
    d = docx.Document(io.BytesIO(buf.getvalue()))
    return "\n".join(p.text for p in d.paragraphs)


def _pptx_text(buf):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(buf.getvalue()))
    return "\n".join(shape.text_frame.text for slide in prs.slides
                     for shape in slide.shapes if shape.has_text_frame)


def test_note_appears_only_for_agent_read_quotes():
    from export_docx import build_proposal_docx
    from export_pptx import generate_proposal
    summary, rec, projection, source_perf = _compute_driven_rec()
    assert NOTE not in _docx_text(build_proposal_docx(summary, rec, projection, source_perf))
    assert NOTE not in _pptx_text(generate_proposal(summary, rec, projection, source_perf))

    flagged = dict(rec, export_override={"bom_agent_read": True})
    assert NOTE in _docx_text(build_proposal_docx(summary, flagged, projection, source_perf))
    assert NOTE in _pptx_text(generate_proposal(summary, flagged, projection, source_perf))
