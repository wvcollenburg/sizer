"""Vendor choice for Validated (software-only) sizing.

A Validated recommendation is built on a vendor's server, so the engine sizes
against ONE vendor's HCL platforms: a sizer model without an active platform
for that vendor drops out, and the recommendation is named after the vendor
chassis instead of "Validated – based off <SC model>". The SC model is
only a footnote on the card ("HCxxxx equivalent"), never a name: a chassis is the product, so its SC variants (all-flash, hybrid, ...)
are listed, targeted and de-duplicated as one. Certified sizing is untouched.

Run: .venv/bin/python -m pytest tests/test_vendor_choice.py -q
"""
import os
import sys

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("ENABLE_SCHEDULER", "0")
os.environ.setdefault("SECRET_KEY", "test-secret")
os.environ.pop("SESSION_COOKIE_SECURE", None)
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

import pytest  # noqa: E402
import app as appmod  # noqa: E402
import orm_models as om  # noqa: E402
import hcl_vendor  # noqa: E402
from database import db  # noqa: E402
from extensions import limiter  # noqa: E402
from hcl_models import HclPlatform, STATUS_DELISTED  # noqa: E402
from tunables import DEFAULTS  # noqa: E402
from recommend import generate_recommendations  # noqa: E402
from fingerprint import catalog_digest  # noqa: E402


@pytest.fixture()
def app():
    application = appmod.app
    application.config["TESTING"] = True
    application.config["RATELIMIT_ENABLED"] = False
    limiter.enabled = False
    with application.app_context():
        db.drop_all()
        db.create_all()
        _seed()
    return application


SC_NAMES = ("HC1450", "HC3450F", "HE155", "HC5250D", "HC1650D")


def _model(name, cpu, nvme, nic):
    m = om.Model(name=name, status="Active", category="3XXX Core", form_factor="1U",
                 chassis="Catalog chassis", min_nodes=2, cost_tier=5.0,
                 validated_only=False)
    db.session.add(m)
    db.session.flush()
    db.session.add(om.ModelCpuOption(model_id=m.id, cpu_id=cpu.id, quantity=2))
    db.session.add(om.RamOption(model_id=m.id, size_gb=512))
    db.session.add(om.ModelNicOption(model_id=m.id, nic_id=nic.id, quantity=1))
    sc = om.StorageConfig(model_id=m.id, storage_type="nvme_only", drives_per_node=6)
    db.session.add(sc)
    db.session.flush()
    db.session.add(om.StorageConfigDrive(storage_config_id=sc.id, drive_id=nvme.id))


def _seed():
    for k, v in DEFAULTS.items():
        db.session.add(om.SizingSetting(key=k, value=float(v)))
    cpu = om.CpuCatalog(description="Xeon Test 32C", cores=32, threads=64, ghz=2.5)
    nvme = om.DriveCatalog(drive_type="NVMe", size_tb=8.0)
    nic = om.NicCatalog(description="Test 25GbE", ports=2, speed="25GbE")
    db.session.add_all([cpu, nvme, nic])
    db.session.flush()
    # HC1450: on Lenovo and Supermicro. HC3450F: a second SC model on the same
    # Lenovo SR630V3 (spelled with a space on its HCL card). HE155-1: the
    # sizer's split of the HCL's HE155. HC5250D: Dell lists it as HC5250D-V.
    # HC1650D: on no vendor.
    for name in ("HC1450", "HC3450F", "HE155-1", "HC5250D", "HC1650D"):
        _model(name, cpu, nvme, nic)
    db.session.add_all([
        HclPlatform(brand="lenovo", sc_model="HC1450", server="ThinkSystem SR630V3"),
        HclPlatform(brand="lenovo", sc_model="HC3450F", server="ThinkSystem SR630 V3"),
        HclPlatform(brand="lenovo", sc_model="HC3450DF", server="ThinkSystem SR630V3"),
        HclPlatform(brand="supermicro", sc_model="HC1450", server="SuperServer SYS-511E-WR"),
        HclPlatform(brand="lenovo", sc_model="HE155", server="ThinkCentre M70q Tiny Gen6"),
        HclPlatform(brand="lenovo", sc_model="HC1650", server="ThinkSystem SR630V4"),
        HclPlatform(brand="dell", sc_model="HC5250D-V", server="PowerEdge R740XD"),
        HclPlatform(brand="hpe", sc_model="HC1400", server="ProLiant DL320 Gen11",
                    status=STATUS_DELISTED),
    ])
    db.session.commit()


SUMMARY = {
    "total_vcpus": 64, "total_vm_provisioned_memory_gb": 200,
    "datastore_used_tb": 5, "nic_speed_mbps": 25000, "total_host_ghz": 400,
    "peak_cpu_ghz": 200, "max_vm_ram_gb": 0, "max_vm_cores": 0,
    "active_vms": 10, "p95_iops": 0, "total_avg_iops": 0,
}


def _rec(app, **kwargs):
    with app.app_context():
        return generate_recommendations(dict(SUMMARY), growth_pct=0, snapshot_pct=0,
                                        years=1, max_day_one_storage_pct=100,
                                        max_day_one_ram_pct=100, **kwargs)


def _chassis(result):
    return {r.get("vendor_chassis") for r in result["recommendations"]}


def _visible_text(rec):
    """Every recommendation field a card or an export prints as a name."""
    return " ".join(str(rec.get(k) or "") for k in ("vendor_chassis", "chassis", "category"))


def test_vendors_list_active_brands_default_first(app):
    with app.app_context():
        vendors = hcl_vendor.list_vendors()
    # HPE's only platform is delisted, so HPE is not a vendor to size on.
    assert [v["brand"] for v in vendors] == ["lenovo", "dell", "supermicro"]
    assert vendors[0]["label"] == "Lenovo"


def test_validated_names_recommendations_by_vendor_chassis(app):
    out = _rec(app, sizing_mode="validated", vendor="lenovo")
    assert out["vendor"] == "lenovo"
    # HC5250D is Dell-only and HC1650D is on no vendor (the HCL's HC1650 is a
    # different model), so neither is a Lenovo Validated box. HC1450 and
    # HC3450F are one chassis with one label despite the HCL's two spellings.
    assert _chassis(out) == {"Lenovo ThinkSystem SR630V3", "Lenovo ThinkCentre M70q Tiny Gen6"}
    for rec in out["recommendations"]:
        assert rec["vendor"] == "lenovo"
        assert rec["chassis"] == rec["vendor_chassis"]
        assert rec["category"] == "1U"
        assert not any(sc in _visible_text(rec) for sc in SC_NAMES), rec
    rec = next(r for r in out["recommendations"] if r["model"] in ("HC1450", "HC3450F"))
    assert rec["refs"]["hcl_platform"] == "lenovo/" + rec["model"]


def test_one_chassis_is_one_product(app):
    """Vendors don't split a server into all-flash/hybrid SC models, so neither
    does Validated sizing: one result per chassis per node count."""
    recs = _rec(app, sizing_mode="validated", vendor="lenovo")["recommendations"]
    seen = [(r["vendor_chassis"], r["node_count"]) for r in recs]
    assert len(seen) == len(set(seen))
    # Certified keeps both SC models apart at the same node count.
    certified = _rec(app, sizing_mode="certified")["recommendations"]
    counts = {r["node_count"] for r in certified if r["model"] == "HC1450"} & \
        {r["node_count"] for r in certified if r["model"] == "HC3450F"}
    assert counts, "fixture should give both SC models a shared node count"


def test_target_a_chassis_spans_its_sc_models(app):
    both = _rec(app, sizing_mode="validated", vendor="lenovo", target_model="lenovo/sr630v3")
    assert _chassis(both) == {"Lenovo ThinkSystem SR630V3"}
    assert {r["model"] for r in both["recommendations"]} <= {"HC1450", "HC3450F"}
    # A sizing saved with an SC model target resolves to that model's chassis.
    legacy = _rec(app, sizing_mode="validated", vendor="lenovo", target_model="HC3450F")
    assert _chassis(legacy) == {"Lenovo ThinkSystem SR630V3"}


def test_other_vendor_changes_the_set_and_the_label(app):
    assert _chassis(_rec(app, sizing_mode="validated", vendor="supermicro")) == {
        "Supermicro SuperServer SYS-511E-WR"}
    # Dell's "-V" suffix still matches the sizer's HC5250D.
    assert _chassis(_rec(app, sizing_mode="validated", vendor="dell")) == {
        "Dell PowerEdge R740XD"}


def test_missing_or_unlisted_vendor_falls_back_to_default(app):
    for vendor in (None, "", "hpe", "nonsense"):
        out = _rec(app, sizing_mode="validated", vendor=vendor)
        assert out["vendor"] == "lenovo", vendor
        assert "Lenovo ThinkSystem SR630V3" in _chassis(out)


def test_certified_ignores_vendor(app):
    out = _rec(app, sizing_mode="certified", vendor="dell")
    assert out["vendor"] is None
    models = {r["model"] for r in out["recommendations"]}
    assert models == {"HC1450", "HC3450F", "HE155-1", "HC5250D", "HC1650D"}
    assert all("vendor_chassis" not in r for r in out["recommendations"])
    assert all("hcl_platform" not in r["refs"] for r in out["recommendations"])


def test_target_model_not_built_by_vendor_warns(app):
    out = _rec(app, sizing_mode="validated", vendor="dell", target_model="lenovo/sr630v3")
    assert out["recommendations"] == []
    assert any("not listed on the HCL for Dell" in w for w in out["warnings"])


def test_empty_hcl_catalog_explains_no_validated_results(app):
    with app.app_context():
        HclPlatform.query.delete()
        db.session.commit()
    out = _rec(app, sizing_mode="validated", vendor="lenovo")
    assert out["recommendations"] == []
    assert out["vendor"] is None
    assert any("HCL scrape" in w for w in out["warnings"])


def test_sc_model_key_keeps_real_model_letters():
    assert hcl_vendor.sc_model_key("HE155-2") == "HE155"
    assert hcl_vendor.sc_model_key("hc5250d-v") == "HC5250D"
    assert hcl_vendor.sc_model_key("HC1650D") != hcl_vendor.sc_model_key("HC1650")


def test_display_model_prefers_chassis():
    assert hcl_vendor.rec_display_model({"model": "HC1450"}) == "HC1450"
    assert hcl_vendor.rec_display_model(
        {"model": "HC1450", "vendor_chassis": "Lenovo ThinkSystem SR630V3"}
    ) == "Lenovo ThinkSystem SR630V3"


def _signed_in(app, email="sa@partnerco.example"):
    c = app.test_client()
    creds = {"email": email, "password": "Abcdef1!xy"}
    c.post("/api/auth/signup", json=dict(creds, accept_privacy=True))
    c.post("/api/auth/login", json=creds)
    return c


def test_models_picker_follows_vendor(app):
    c = _signed_in(app)
    certified = c.get("/api/models?mode=appliance&status=active&sizing=certified").get_json()
    assert set(certified) == {"HC1450", "HC3450F", "HE155-1", "HC5250D", "HC1650D"}
    dell = c.get("/api/models?mode=appliance&status=active&sizing=validated&vendor=dell").get_json()
    assert dell == {"dell/r740xd": {"category": "1U", "status": "Active",
                                    "vendor_chassis": "Dell PowerEdge R740XD"}}
    # Two SC models on one chassis are one picker entry, with no SC name in it.
    lenovo = c.get("/api/models?mode=appliance&status=active&sizing=validated&vendor=lenovo").get_json()
    assert set(lenovo) == {"lenovo/sr630v3", "lenovo/m70qtinygen6"}
    assert not any(sc in str(lenovo) for sc in SC_NAMES)


def test_sizing_page_renders_vendor_options(app):
    html = app.test_client().get("/").get_data(as_text=True)
    assert 'id="sizing-vendor"' in html and 'id="dr-vendor"' in html
    assert '<option value="lenovo">Lenovo</option>' in html
    assert '<option value="hpe">' not in html


def test_fingerprint_tracks_the_hcl_platform(app):
    with app.app_context():
        refs = _rec(app, sizing_mode="validated", vendor="lenovo")
        rec = next(r for r in refs["recommendations"] if r["model"] == "HC1450")
        before = catalog_digest(rec["refs"])
        # Certified refs for the same model carry no platform and keep their digest.
        plain = dict(rec["refs"])
        plain.pop("hcl_platform")
        certified_before = catalog_digest(plain)

        p = HclPlatform.query.filter_by(brand="lenovo", sc_model="HC1450").one()
        p.status = STATUS_DELISTED
        db.session.commit()
        assert catalog_digest(rec["refs"]) != before
        assert catalog_digest(plain) == certified_before


def _document_text(path):
    if path.endswith(".pptx"):
        from pptx import Presentation
        return "\n".join(sh.text_frame.text for slide in Presentation(path).slides
                         for sh in slide.shapes if sh.has_text_frame)
    from docx import Document
    doc = Document(path)
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


@pytest.mark.parametrize("fmt", ["pptx", "docx"])
def test_exports_and_comparison_name_the_vendor_chassis(app, fmt):
    import export_worker
    from project_models import ExportJob

    summary = dict(SUMMARY, total_vms=12, total_host_cores=64, hosts=3,
                   total_host_ram_gb=768, datastore_total_tb=20)
    with app.app_context():
        result = generate_recommendations(summary, 3.0, sizing_mode="validated",
                                          vendor="supermicro")
    rec, projection = result["recommendations"][0], result["projection"]
    assert rec["vendor_chassis"] == "Supermicro SuperServer SYS-511E-WR"

    # Editable PPTX/DOCX are Scale-only, so export as a Scale user.
    c = _signed_in(app, "sa@scalecomputing.com")
    project = c.post("/api/projects/", json={"name": "Acme"}).get_json()
    row = c.post("/api/configs/", json={"name": "Site A", "payload": {"mode": "import"},
                                        "project_id": project["id"]}).get_json()
    assert c.put(f"/api/sizings/{row['id']}/result", json={"clusters": [{
        "name": "Prod", "summary": summary, "projection": projection,
        "recommendation": rec, "source_perf": None, "replicates_to": "",
        "refs": rec["refs"]}]}).status_code == 200

    compare = c.post(f"/api/projects/{project['id']}/compare",
                     json={"sizing_ids": [row["id"]]})
    assert compare.status_code == 200, compare.get_data(as_text=True)
    assert "Supermicro SuperServer SYS-511E-WR" in compare.get_data(as_text=True)
    assert "HC1450" not in compare.get_json().get("sizings", [{}])[0].get("totals", {}).get("model", "")

    queued = c.post(f"/api/projects/{project['id']}/export",
                    json={"format": fmt, "sizing_ids": [row["id"]]})
    assert queued.status_code == 202, queued.get_data(as_text=True)
    with app.app_context():
        export_worker.run_job(export_worker.claim_next_job(), app)
        job = db.session.get(ExportJob, queued.get_json()["id"])
        assert job.status == "done", job.error
        text = _document_text(job.artifact_path)
    assert "Supermicro SuperServer SYS-511E-WR" in text
    assert "based off" not in text
    assert not any(sc in text for sc in SC_NAMES), [sc for sc in SC_NAMES if sc in text]
