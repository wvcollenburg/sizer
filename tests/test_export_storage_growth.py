"""Exports with separate storage growth, and a DR target that follows its
sources (plan A, phase 3).

The pptx projection slide and the docx growth sentences use the split keys
when the toggle is on, the old keys when it is off, and still render an old
projection stored without the new keys. A dedicated DR target whose demand was
pre-grown shows its capacity planning with one sentence naming each source's
growth; one sized before that keeps the old skip.

Run:  .venv/bin/python -m pytest tests/test_export_storage_growth.py -q
"""
import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

import pytest  # noqa: E402
from docx import Document  # noqa: E402
from pptx import Presentation  # noqa: E402

import test_perf_sizing_e2e as e2e  # noqa: E402
from i18n import translator  # noqa: E402

HEADING = translator("en")("export.docx.capacity_planning", years=5)


@pytest.fixture(scope="module")
def app():
    app = e2e._build_app()
    from database import db
    with app.app_context():
        db.create_all()
        e2e._seed_catalog()
    return app


def _sized(app, **kw):
    from recommend import generate_recommendations
    kw.setdefault("growth_pct", 10)
    kw.setdefault("snapshot_pct", 20)
    kw.setdefault("years", 5)
    with app.app_context():
        res = generate_recommendations(e2e._summary(), max_day_one_storage_pct=100,
                                       max_day_one_ram_pct=100, **kw)
    return res["recommendations"][0], res["projection"]


def _pptx_text(buf):
    out = []
    for slide in Presentation(io.BytesIO(buf.getvalue())).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.extend(c.text for c in row.cells)
    return "\n".join(out)


def _docx_text(buf):
    doc = Document(io.BytesIO(buf.getvalue()))
    out = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            out.extend(c.text for c in row.cells)
    return "\n".join(out)


def _both(summary, rec, proj):
    from export_docx import build_proposal_docx
    from export_pptx import generate_proposal
    return (_pptx_text(generate_proposal(summary, rec, proj)),
            _docx_text(build_proposal_docx(summary, rec, proj)))


def test_toggle_off_uses_the_single_rate_copy(app):
    rec, proj = _sized(app)
    ppt, doc = _both(e2e._summary(), rec, proj)
    assert "10% per year" in ppt
    assert "Compute 10%" not in ppt
    assert "10% YoY growth, 20% snapshot overhead" in doc
    assert "for storage" not in doc


def test_toggle_on_uses_the_split_copy(app):
    rec, proj = _sized(app, growth_pct=0, storage_growth_pct=20)
    ppt, doc = _both(e2e._summary(), rec, proj)
    assert "Compute 0% · Storage 20% per year" in ppt
    assert "Compute 1.0x · Storage 2.488x over 5 years" in ppt
    assert "0% compute / 20% storage YoY growth" in ppt
    assert "0% YoY compute growth, 20% YoY storage growth" in doc
    assert "at 0% annually for compute and 20% for storage" in doc


def test_old_projection_without_new_keys_still_renders(app):
    rec, proj = _sized(app)
    for k in ("storage_growth_pct", "storage_growth_factor", "separate_storage_growth"):
        proj.pop(k)
    ppt, doc = _both(e2e._summary(), rec, proj)
    assert "10% per year" in ppt
    assert "10% YoY growth" in doc


def test_split_copy_in_a_translation(app):
    from export_pptx import generate_proposal
    rec, proj = _sized(app, growth_pct=5, storage_growth_pct=25)
    ppt = _pptx_text(generate_proposal(e2e._summary(), rec, proj, lang="de"))
    assert "Compute 5 % · Speicher 25 % pro Jahr" in ppt


# ── a dedicated DR target ─────────────────────────────────────────────────────

def _dr_cluster(app, pregrown=True):
    import projects
    reserve = {"vcpus": 64, "ram_gb": 200, "storage_tb": 5}
    summary = projects._dr_demand_summary(reserve)
    kw = {}
    if pregrown:
        kw["pregrown"] = {"vcpus": 90, "ram_gb": 280, "storage_tb": 9,
                          "snapshot_tb": 1.5, "years": 5}
    from recommend import generate_recommendations
    with app.app_context():
        res = generate_recommendations(summary, max_day_one_storage_pct=100,
                                       max_day_one_ram_pct=100, **kw)
    proj = res["projection"]
    if pregrown:
        proj["dr_sources"] = [
            {"sizing_name": "HQ", "years": 3, "growth_pct": 5,
             "storage_growth_pct": 5, "snapshot_pct": 10},
            {"sizing_name": "Branch", "years": 5, "growth_pct": 15,
             "storage_growth_pct": 30, "snapshot_pct": 20}]
    return {"name": "DR site", "summary": summary,
            "recommendation": res["recommendations"][0], "projection": proj,
            "source_perf": None, "replicates_to": ""}


def _prod_cluster(app):
    rec, proj = _sized(app)
    return {"name": "HQ", "summary": e2e._summary(), "recommendation": rec,
            "projection": proj, "source_perf": None, "replicates_to": "DR site"}


def test_dr_bundle_docx_names_each_source_growth(app):
    from export_docx import build_bundle_proposal_docx
    doc = _docx_text(build_bundle_proposal_docx([_prod_cluster(app), _dr_cluster(app)]))
    assert ("The replicas are grown at each protected site's own sizing options: "
            "HQ at 5% per year over 3 years (storage 5%, snapshot 10%); "
            "Branch at 15% per year over 5 years (storage 30%, snapshot 20%).") in doc
    # Capacity planning shows for both sites now.
    assert doc.count(HEADING) == 2


def test_old_dr_bundle_keeps_the_skip(app):
    from export_docx import build_bundle_proposal_docx
    doc = _docx_text(build_bundle_proposal_docx(
        [_prod_cluster(app), _dr_cluster(app, pregrown=False)]))
    assert "The replicas are grown" not in doc
    assert doc.count(HEADING) == 1


def test_dr_bundle_pptx_projection_reads_as_following_sources(app):
    from export_pptx import generate_bundle_proposal
    ppt = _pptx_text(generate_bundle_proposal([_prod_cluster(app), _dr_cluster(app)]))
    assert "Growth follows each protected site's own sizing" in ppt
