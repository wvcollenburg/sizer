"""Export customization — the chassis (and hardware) the exports name.

A Validated sizing is sized on one vendor chassis, but the partner quoting it
may sell another one: we size a Dell R660, the VAR BOM-checks an R670, and the
proposal has to say R670. Three sources merged PER FIELD:

    manual override  >  linked BOM check  >  the recommendation

These tests pin the owner's decisions of 2026-09-16 and 2026-09-17:
  * only Validated recommendations are ever renamed;
  * a passing BOM check applies by itself, a failing one never does;
  * per-node hardware follows the source, and every figure that follows from
    it follows too: totals, N-1, usable storage, the vCPU ratio, the compute
    floor, the licence requirement and the utilization bars (demand kept,
    capacity replaced, over 100 % NOT clamped);
  * manual disks are structured and use the engine's own storage maths;
  * the stored snapshot is never mutated.

Run: .venv/bin/python -m pytest tests/test_export_override.py -q
"""
import os
import sys

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("ENABLE_SCHEDULER", "0")
os.environ.setdefault("SECRET_KEY", "test-secret")
os.environ.pop("SESSION_COOKIE_SECURE", None)
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))

import pytest  # noqa: E402

import app as appmod  # noqa: E402
import export_override as eo  # noqa: E402
from auth_models import Configuration, Tenant, User  # noqa: E402
from bom_models import BomCheck  # noqa: E402
from database import db  # noqa: E402
from extensions import limiter  # noqa: E402
from project_models import Project  # noqa: E402


# ── fixtures ─────────────────────────────────────────────────────────────────

def _rec(validated=True):
    """A stored Validated recommendation, trimmed to the fields that matter."""
    return {
        "validated": validated,
        "model": "HC3450DF",
        "vendor": "dell",
        "vendor_chassis": "Dell PowerEdge R660" if validated else None,
        "chassis": "Dell PowerEdge R660",
        "category": "1U",
        "form_factor": "1U",
        "cpu": "2 x Intel Xeon Gold 6526Y",
        "cores_per_node": 32,
        "usable_cores_per_node": 30,
        "threads_per_node": 64,
        "ghz": 2.8,
        "ram_per_node_gb": 512,
        "usable_ram_per_node_gb": 480,
        "node_count": 4,
        "hci_node_count": 4,
        "cluster_layout": [4],
        "totals": {"cores": 120, "threads": 256, "total_ghz": 358.4,
                   "ram_gb": 1920.0, "raw_storage_tb": 61.4,
                   "usable_storage_tb": 24.0},
        "n_minus_1": {"cores": 90, "threads": 192, "total_ghz": 268.8,
                      "ram_gb": 1440.0, "usable_storage_tb": 24.0},
        "storage_config": {"desc": "4 x 3.84TB NVMe", "raw_per_node": 15.36,
                           "biggest_disk": 3.84, "drive_counts": {"NVMe": 4}},
        "cpu_perf_index": 400.0,
        "vcpu_ratio": 3.75,
        "vcpu_ratio_degraded": 3.75,
        "sized_full_cluster": False,
        "compute_floor": {"coverage_pct": 110.0, "ghz_pct": 120.0,
                          "perf_pct": 100.0, "balance": 0.5,
                          "source_cpu_util_pct": 40.0},
        "iops": {"per_node": 100000, "total": 400000, "n_minus_1": 300000},
        # Demand in real units (`abs`) and percentages of the sized capacity,
        # the shape recommend.py stores.
        "utilization": {
            "cpu": {"current": 55, "total": 70, "replication": 0, "snapshot": 0,
                    "ha_reserve": 25,
                    "abs": {"current": 66, "total": 84, "snapshot": 0,
                            "capacity": 120, "unit": "cores"}},
            "ram": {"current": 50, "total": 65, "replication": 0, "snapshot": 0,
                    "ha_reserve": 25,
                    "abs": {"current": 960.0, "total": 1248.0, "snapshot": 0,
                            "capacity": 1920.0, "unit": "GB"}},
            "storage": {"current": 40, "total": 60, "replication": 0,
                        "snapshot": 10, "ha_reserve": 0,
                        "abs": {"current": 9.6, "total": 14.4, "snapshot": 2.4,
                                "capacity": 24.0, "unit": "TB"}},
        },
        "licensing": {"band": "49-64"},
    }


def _snapshot(rec=None):
    # summary/projection must be non-empty: the export worker only renders a
    # cluster that carries all three blocks.
    return {"clusters": [{"name": "Prod",
                          "summary": {"vm_count": 40},
                          "projection": {"years": 3},
                          "recommendation": rec or _rec()}],
            "totals": {}, "tunables": "x"}


def _bom_result(server="PowerEdge R670", brand="dell", nodes=5, cores=48,
                ram=768, threads=96, with_cluster=True, drives=None):
    """The shape bom/check.py stores: technical config_results + the fit block
    whose `cluster` is bom/fit.py's derived cluster summary."""
    result = {
        "technical": {"verdict": "PASS", "config_results": [{
            "config_name": "R670",
            "verdict": "PASS",
            "platform": {"brand": brand, "server": server, "form_factor": "2U",
                         "sc_models": ["HC3450DF"], "platform_ids": [1]},
            "node_count": nodes,
        }]},
        "fit": {"verdict": "match", "dimensions": [], "notes": []},
    }
    if with_cluster:
        result["fit"]["cluster"] = {
            "node_count": nodes, "hci_node_count": nodes, "so_count": 0,
            "cluster_layout": [nodes], "cpu": "2 x Intel Xeon Gold 6542Y",
            "cores_per_node": cores, "threads_per_node": threads,
            "sockets_per_node": 2,
            "ram_per_node_gb": ram, "usable_ram_per_node_gb": ram - 32,
            "cores_full": (cores - 2) * nodes, "cores_n1": (cores - 2) * (nodes - 1),
            "ram_full": (ram - 32) * nodes, "ram_n1": (ram - 32) * (nodes - 1),
            "ghz_full": 400.0, "ghz_n1": 320.0,
            "raw_storage_tb": 92.2, "usable_storage_tb": 36.0,
            "drive_counts": {"NVMe": 6}, "bays": 10,
            "storage_category": "flash", "nic_gbe": 25, "nic_ports": 2,
            "perf_full": 500.0 * nodes, "perf_n1": 500.0 * (nodes - 1),
        }
        if drives is not None:
            # bom/fit.py stores per-node disk sizes since 2026-09-17; the
            # default fixture above is an older check that only counted them.
            result["fit"]["cluster"]["drives"] = drives
    return result


@pytest.fixture()
def app():
    application = appmod.app
    application.config["TESTING"] = True
    application.config["RATELIMIT_ENABLED"] = False
    limiter.enabled = False
    with application.app_context():
        db.drop_all()
        db.create_all()
        yield application
        db.session.remove()


@pytest.fixture()
def sizing(app):
    tenant = Tenant(domain="example.com")
    db.session.add(tenant)
    db.session.flush()
    user = User(email="sa@example.com", password_hash="x", tenant_id=tenant.id,
                is_verified=True)
    db.session.add(user)
    db.session.flush()
    project = Project(name="Acme", owner_id=user.id, tenant_id=tenant.id,
                      code="PRJ1")
    db.session.add(project)
    db.session.flush()
    config = Configuration(code="SZ1", name="Option 1", owner_id=user.id,
                           tenant_id=tenant.id, project_id=project.id,
                           payload={}, result_snapshot=_snapshot())
    db.session.add(config)
    db.session.commit()
    return config


def _add_check(sizing, name="VAR quote", result=None, technical="PASS",
               fit_verdict="match", vendor="dell"):
    check = BomCheck(project_id=sizing.project_id, configuration_id=sizing.id,
                     user_id=sizing.owner_id, tenant_id=sizing.tenant_id,
                     name=name, vendor=vendor,
                     normalized={"vendor": vendor, "configs": []},
                     result=result if result is not None else _bom_result(),
                     technical_verdict=technical, fit_verdict=fit_verdict)
    db.session.add(check)
    db.session.commit()
    return check


# ── the stored setting ───────────────────────────────────────────────────────

def test_clean_manual_keeps_known_fields_and_drops_the_rest():
    cleaned = eo.clean_manual({"chassis": "  Dell PowerEdge R670  ",
                               "cores_per_node": "48", "node_count": 0,
                               "ram_per_node_gb": "", "cpu": "   ",
                               "score": 9999})
    assert cleaned == {"chassis": "Dell PowerEdge R670", "cores_per_node": 48}


def test_default_setting_stores_as_nothing():
    assert eo.is_default(None)
    assert eo.is_default({"bom": "auto", "manual": {}})
    assert not eo.is_default({"bom": "none", "manual": {}})
    assert not eo.is_default({"bom": "auto", "manual": {"chassis": "R670"}})


# ── which BOM check drives it ────────────────────────────────────────────────

def test_a_passing_bom_check_applies_by_itself(sizing):
    _add_check(sizing)
    badge = eo.badge_for(sizing)
    assert badge["chassis"] == "Dell PowerEdge R670"
    assert badge["bom_auto"] is True


def test_a_failing_check_is_never_picked_automatically(sizing):
    _add_check(sizing, name="Bad quote", technical="FAIL")
    assert eo.badge_for(sizing) is None


def test_a_check_under_the_requirement_is_never_picked_automatically(sizing):
    _add_check(sizing, name="Too small", fit_verdict="smaller")
    assert eo.badge_for(sizing) is None


def test_a_failing_check_can_still_be_chosen_by_hand(sizing):
    check = _add_check(sizing, name="Bad quote", technical="FAIL")
    sizing.export_override = {"bom": check.id, "manual": {}}
    badge = eo.badge_for(sizing)
    assert badge["chassis"] == "Dell PowerEdge R670"
    assert badge["bom_auto"] is False


def test_bom_none_keeps_the_recommendation(sizing):
    _add_check(sizing)
    sizing.export_override = {"bom": "none", "manual": {}}
    assert eo.badge_for(sizing) is None


def test_a_check_without_a_usable_cluster_is_not_offered(sizing):
    # Old/degraded results carry a fit block with no cluster at all.
    _add_check(sizing, result=_bom_result(with_cluster=False))
    assert eo.candidate_checks(sizing) == []
    assert eo.badge_for(sizing) is None


def test_a_deleted_chosen_check_falls_back_to_the_suggestion(sizing):
    good = _add_check(sizing, name="Current")
    sizing.export_override = {"bom": good.id + 999, "manual": {}}
    badge = eo.badge_for(sizing)
    assert badge["bom_check_id"] == good.id
    assert badge["bom_auto"] is True


# ── applying it to a recommendation ──────────────────────────────────────────

def test_certified_recommendations_are_never_renamed(sizing):
    _add_check(sizing)
    override = eo.resolve(sizing)
    certified = _rec(validated=False)
    certified.pop("vendor_chassis")
    assert eo.apply_to_rec(certified, override) is certified


def test_bom_supplies_chassis_and_per_node_hardware(sizing):
    _add_check(sizing)
    override = eo.resolve(sizing)
    out = eo.apply_to_rec(_rec(), override)
    assert out["vendor_chassis"] == "Dell PowerEdge R670"
    assert out["chassis"] == "Dell PowerEdge R670"
    assert out["form_factor"] == "2U"
    assert out["cpu"] == "2 x Intel Xeon Gold 6542Y"
    assert out["cores_per_node"] == 48
    assert out["threads_per_node"] == 96
    assert out["ram_per_node_gb"] == 768
    assert out["node_count"] == 5
    assert "6 x NVMe" in out["storage_config"]["desc"]


def test_the_sc_model_stays_as_catalog_identity(sizing):
    _add_check(sizing)
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    # refs/fingerprint/BOM-fit all ride on `model`; only the display name moves.
    assert out["model"] == "HC3450DF"


def test_cluster_totals_follow_the_quoted_hardware(sizing):
    _add_check(sizing)
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    cluster = _bom_result()["fit"]["cluster"]
    assert out["totals"]["cores"] == cluster["cores_full"]
    assert out["totals"]["ram_gb"] == cluster["ram_full"]
    assert out["totals"]["usable_storage_tb"] == cluster["usable_storage_tb"]
    assert out["n_minus_1"]["cores"] == cluster["cores_n1"]


# ── bars, ratio, compute floor, licensing follow the quote (2026-09-17) ─────

def test_bars_keep_the_demand_and_follow_the_quoted_capacity(sizing):
    """The workload is the same, so demand (abs) is kept; only capacity moves.
    Quoted: 5 x 48C (46 usable) / 768 GB (736 usable) / 36 TB usable."""
    _add_check(sizing)
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    cpu, ram, st = (out["utilization"][k] for k in ("cpu", "ram", "storage"))
    assert cpu["abs"]["current"] == 66 and cpu["abs"]["total"] == 84   # demand kept
    assert cpu["abs"]["capacity"] == 230
    assert (cpu["current"], cpu["total"]) == (29, 37)                 # 66/230, 84/230
    assert cpu["ha_reserve"] == 20                                    # (230-184)/230
    assert ram["abs"]["capacity"] == 3680.0
    assert (ram["current"], ram["total"], ram["ha_reserve"]) == (26, 34, 20)
    assert st["abs"]["capacity"] == 36.0
    assert (st["current"], st["total"], st["snapshot"]) == (27, 40, 7)
    assert st["ha_reserve"] == 0


def test_an_undersized_quote_reads_over_100_percent_not_clamped(sizing):
    """A 2-node 16-core quote cannot carry 84 cores of demand. Clamped, it
    would draw as 'exactly full' - the opposite of the truth."""
    _add_check(sizing, result=_bom_result(nodes=2, cores=16, threads=32, ram=256))
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    cpu = out["utilization"]["cpu"]
    assert cpu["abs"]["capacity"] == 14 * 2
    assert cpu["total"] == 300 and cpu["current"] > 100


def test_full_cluster_sizing_releases_the_cpu_ha_band_only(sizing):
    rec = _rec()
    rec["sized_full_cluster"] = True
    _add_check(sizing)
    out = eo.apply_to_rec(rec, eo.resolve(sizing))
    assert out["utilization"]["cpu"]["ha_reserve"] == 0
    assert out["utilization"]["ram"]["ha_reserve"] == 20


def test_the_vcpu_ratio_follows_the_quoted_cores(sizing):
    """Same vCPUs over 184 N-1 cores instead of 90."""
    _add_check(sizing)
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["vcpu_ratio"] == round(3.75 * 90 / 184, 2)
    assert out["vcpu_ratio_degraded"] == round(3.75 * 90 / 184, 2)


def test_the_compute_floor_follows_a_bom_cpu(sizing):
    """Coverage scales with (per-node supply x compute nodes). Sized: 2.8 GHz x
    32C = 89.6 GHz and 400 perf per node over 3 N-1 nodes. Quoted: 80 GHz and
    500 perf per node over 4 N-1 nodes."""
    _add_check(sizing)
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    cf = out["compute_floor"]
    ghz = round(120.0 * (80.0 * 4) / (89.6 * 3), 1)
    perf = round(100.0 * (500.0 * 4) / (400.0 * 3), 1)
    assert cf["ghz_pct"] == ghz and cf["perf_pct"] == perf
    assert cf["coverage_pct"] == round(0.5 * ghz + 0.5 * perf, 1)
    assert out["cpu_perf_index"] == 500.0


def test_the_compute_floor_is_dropped_for_a_cpu_typed_by_hand(sizing):
    """No catalog behind a typed CPU: its clock and benchmark are unknown, so
    the coverage line is left out rather than guessed (owner decision)."""
    sizing.export_override = {"bom": "none", "manual": {
        "cpu": "2 x Some Future CPU", "cores_per_node": 64}}
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["compute_floor"] is None
    assert out["cpu_perf_index"] is None


def test_the_compute_floor_rescales_exactly_for_a_node_count_change(sizing):
    """Same CPU, 6 nodes instead of 4: 5 N-1 compute nodes instead of 3."""
    sizing.export_override = {"bom": "none", "manual": {"node_count": 6}}
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    cf = out["compute_floor"]
    assert cf["ghz_pct"] == round(120.0 * 5 / 3, 1)
    assert cf["perf_pct"] == round(100.0 * 5 / 3, 1)


def test_licensing_is_recomputed_for_the_quoted_hardware(sizing, monkeypatch):
    import recommend
    seen = {}

    def fake(layout, cores_per_node, ram_gb_per_node, term_years=None, region=None):
        seen.update(layout=layout, cores=cores_per_node, ram=ram_gb_per_node)
        return {"band": "quoted"}

    monkeypatch.setattr(recommend, "license_annotations_for", fake)
    _add_check(sizing)
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["licensing"] == {"band": "quoted"}
    assert seen == {"layout": [5], "cores": 48, "ram": 768}


def test_a_rename_only_override_leaves_every_figure_alone(sizing):
    """Chassis only: nothing about the hardware changed, so the engine's own
    numbers stay byte-for-byte - no rounding drift from a needless rebuild."""
    sizing.export_override = {"bom": "none", "manual": {"chassis": "Dell PowerEdge R670"}}
    rec = _rec()
    out = eo.apply_to_rec(rec, eo.resolve(sizing))
    for key in ("utilization", "totals", "n_minus_1", "compute_floor",
                "licensing", "vcpu_ratio", "iops", "storage_config"):
        assert out[key] == rec[key], key


# ── disks (2026-09-17) ───────────────────────────────────────────────────────

def test_clean_manual_keeps_complete_drive_rows_only():
    cleaned = eo.clean_manual({"drives": [
        {"type": "nvme", "size_tb": "7,68", "count": "4"},
        {"type": "hdd", "size_tb": "", "count": 3},           # incomplete: dropped
        {"type": "tape", "size_tb": 8, "count": 3},           # unknown type
    ]})
    assert cleaned == {"drives": [{"kind": "NVMe", "capacity_tb": 7.68, "qty_per_node": 4}]}


def test_manual_disks_use_the_engines_usable_storage_maths(sizing):
    """Hybrid, 4 nodes: RF2 plus one rebuild disk per cluster, exactly as
    recommend._cluster_usable_storage computes it."""
    from recommend import _cluster_usable_storage
    sizing.export_override = {"bom": "none", "manual": {"drives": [
        {"type": "HDD", "size_tb": 8, "count": 3},
        {"type": "NVMe", "size_tb": 7.68, "count": 1}]}}
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    raw_pn = 3 * 8 + 7.68
    assert out["storage_config"]["desc"] == "3 x 8 TB HDD + 1 x 7.68 TB NVMe"
    assert out["totals"]["raw_storage_tb"] == round(raw_pn * 4, 2)
    assert out["totals"]["usable_storage_tb"] == round(
        _cluster_usable_storage(raw_pn, 8.0, [4]), 2)
    assert out["utilization"]["storage"]["abs"]["capacity"] == out["totals"]["usable_storage_tb"]
    # other disks: the sized IOPS no longer describe the node
    assert out["iops"] is None


def test_a_node_count_change_relays_out_the_sized_disks(sizing):
    """Limit fixed 2026-09-17: a manual node count used to leave storage as
    sized. Same disks per node over 6 nodes is exact, not an estimate."""
    sizing.export_override = {"bom": "none", "manual": {"node_count": 6}}
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["totals"]["usable_storage_tb"] == round((15.36 * 6 - 3.84) / 2, 2)
    assert out["totals"]["raw_storage_tb"] == round(15.36 * 6, 2)
    assert out["iops"]["total"] == 100000 * 6          # same disks: re-multiplied


def test_a_bom_check_names_its_disks(sizing):
    _add_check(sizing, result=_bom_result(drives=[
        {"kind": "NVMe", "capacity_tb": 7.68, "qty_per_node": 6}]))
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["storage_config"]["desc"] == "6 x 7.68 TB NVMe"
    assert out["totals"]["usable_storage_tb"] == round((46.08 * 5 - 7.68) / 2, 2)


def test_an_older_bom_check_without_disk_sizes_keeps_the_count_line(sizing):
    _add_check(sizing)                       # fixture: no `drives` stored
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["storage_config"]["desc"].startswith("6 x NVMe (")
    assert out["totals"]["usable_storage_tb"] == 36.0     # the check's own figure


def test_manual_disks_win_over_the_bom_disks(sizing):
    _add_check(sizing, result=_bom_result(drives=[
        {"kind": "NVMe", "capacity_tb": 7.68, "qty_per_node": 6}]))
    sizing.export_override = {"bom": "auto", "manual": {"drives": [
        {"type": "SSD", "size_tb": 3.84, "count": 8}]}}
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["storage_config"]["desc"] == "8 x 3.84 TB SSD"
    assert out["cores_per_node"] == 48                   # CPU still from the BOM


def test_manual_wins_field_by_field_over_the_bom(sizing):
    _add_check(sizing)
    sizing.export_override = {"bom": "auto",
                              "manual": {"chassis": "Dell PowerEdge R6715"}}
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["vendor_chassis"] == "Dell PowerEdge R6715"   # manual
    assert out["cores_per_node"] == 48                       # still the BOM's
    assert out["ram_per_node_gb"] == 768


def test_manual_only_override_without_any_bom_check(sizing):
    sizing.export_override = {"bom": "auto", "manual": {
        "chassis": "Supermicro AS-1115HS", "cores_per_node": 64,
        "node_count": 3, "ram_per_node_gb": 1024}}
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["vendor_chassis"] == "Supermicro AS-1115HS"
    # The OS overheads are carried across from the sized config rather than
    # re-derived: 32 - 30 cores and 512 - 480 GB.
    assert out["usable_cores_per_node"] == 62
    assert out["usable_ram_per_node_gb"] == 992
    assert out["totals"]["cores"] == 62 * 3
    assert out["totals"]["ram_gb"] == 992 * 3


def test_a_quoted_node_count_splits_hci_and_storage_only(sizing):
    """node_count is the whole cluster: the HCI count is what is left after the
    storage-only nodes, which is what every cluster figure is summed over."""
    rec = _rec()
    rec["storage_only"] = {"count": 2, "cpu": "1 x Xeon 4310"}
    rec["hci_node_count"] = 2
    sizing.export_override = {"bom": "auto", "manual": {"node_count": 6}}
    out = eo.apply_to_rec(rec, eo.resolve(sizing))
    assert out["node_count"] == 6
    assert out["hci_node_count"] == 4


def test_a_bom_cluster_states_its_own_hci_split(sizing):
    _add_check(sizing)
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["node_count"] == 5
    assert out["hci_node_count"] == 5


def test_the_stored_snapshot_is_never_mutated(sizing):
    _add_check(sizing)
    before = sizing.result_snapshot["clusters"][0]["recommendation"]["vendor_chassis"]
    patched = eo.apply_to_snapshot(sizing)
    after = sizing.result_snapshot["clusters"][0]["recommendation"]["vendor_chassis"]
    assert before == after == "Dell PowerEdge R660"
    assert patched["clusters"][0]["recommendation"]["vendor_chassis"] \
        == "Dell PowerEdge R670"


def test_vendor_mismatch_is_followed_and_reported(sizing):
    _add_check(sizing, result=_bom_result(server="ThinkSystem SR630 V3",
                                          brand="lenovo"), vendor="lenovo")
    badge = eo.badge_for(sizing)
    # The BOM is what is being ordered (owner decision), so it wins — but the
    # vendor it came from is on the badge for the warning the UI shows.
    assert badge["chassis"] == "Lenovo ThinkSystem SR630 V3"
    assert badge["vendor"] == "lenovo"


def test_a_server_line_already_naming_its_brand_is_not_doubled(sizing):
    _add_check(sizing, result=_bom_result(server="Dell PowerEdge R670"))
    assert eo.badge_for(sizing)["chassis"] == "Dell PowerEdge R670"


def test_badges_for_resolves_a_whole_project_in_one_pass(sizing):
    _add_check(sizing)
    badges = eo.badges_for([sizing])
    assert badges[sizing.id]["chassis"] == "Dell PowerEdge R670"


# ── the API ──────────────────────────────────────────────────────────────────

@pytest.fixture()
def client(app, sizing):
    with app.test_client() as c:
        with c.session_transaction() as sess:
            sess["user_id"] = sizing.owner_id
            sess["_fresh"] = True
        yield c


def test_get_returns_the_setting_and_the_candidates(client, sizing):
    _add_check(sizing)
    res = client.get(f"/api/sizings/{sizing.id}/export-override")
    assert res.status_code == 200
    body = res.get_json()
    assert body["setting"] == {"bom": "auto", "manual": {}}
    assert body["effective"]["chassis"] == "Dell PowerEdge R670"
    assert len(body["checks"]) == 1
    assert body["checks"][0]["values"]["cores_per_node"] == 48


def test_put_stores_and_clears_the_override(client, sizing):
    res = client.put(f"/api/sizings/{sizing.id}/export-override",
                     json={"bom": "none", "manual": {"chassis": "R670"}})
    assert res.status_code == 200
    assert db.session.get(Configuration, sizing.id).export_override == {
        "bom": "none", "manual": {"chassis": "R670"}}

    res = client.put(f"/api/sizings/{sizing.id}/export-override",
                     json={"bom": "auto", "manual": {}})
    assert res.status_code == 200
    # A default setting stores as NULL, so "has an override" stays a non-null
    # test in queries.
    assert db.session.get(Configuration, sizing.id).export_override is None


def test_the_comparison_screen_shows_the_quoted_chassis(client, sizing):
    """The comparison and compare.xlsx read the same helper as the exports, so
    a chosen option cannot read as one box on screen and another in the file."""
    _add_check(sizing)
    res = client.post(f"/api/projects/{sizing.project_id}/compare",
                      json={"sizing_ids": [sizing.id]})
    assert res.status_code == 200
    rows = res.get_json()["rows"]
    assert rows[0]["totals"]["model"] == "Dell PowerEdge R670"


def test_put_rejects_someone_elses_sizing(app, sizing):
    other = User(email="other@example.com", password_hash="x",
                 tenant_id=sizing.tenant_id, is_verified=True)
    db.session.add(other)
    db.session.commit()
    with appmod.app.test_client() as c:
        with c.session_transaction() as sess:
            sess["user_id"] = other.id
            sess["_fresh"] = True
        res = c.put(f"/api/sizings/{sizing.id}/export-override",
                    json={"bom": "none", "manual": {}})
    assert res.status_code == 403


def test_the_export_worker_renders_the_quoted_chassis(app, sizing):
    """The whole point: what the bundle generators receive.

    sections_for is the single funnel every format goes through, so this is
    also what the PPTX/Word/PDF titles and the per-node table end up naming.
    """
    from project_models import ExportJob
    from export_worker import sections_for

    _add_check(sizing)
    job = ExportJob(user_id=sizing.owner_id, project_id=sizing.project_id,
                    fmt="pptx", sizing_ids=[sizing.id], lang="en")
    db.session.add(job)
    db.session.commit()

    sections, skipped = sections_for(job)
    assert not skipped
    rec = sections[0]["recommendation"]
    assert rec["vendor_chassis"] == "Dell PowerEdge R670"
    assert rec["cores_per_node"] == 48
    # and the sizing itself still holds what was sized
    assert sizing.result_snapshot["clusters"][0]["recommendation"][
        "vendor_chassis"] == "Dell PowerEdge R660"


def test_a_real_proposal_names_the_quoted_chassis(app):
    """End-to-end on REAL engine output: seed a catalog, size a Validated
    cluster, attach a BOM check for another chassis, and read the rendered
    PPTX/DOCX back.

    The hand-built recommendation above is a trimmed fixture; this one is
    whatever the engine actually produces today, which is the only way to know
    the document does not still say the sized chassis somewhere.
    """
    import io

    from docx import Document
    from pptx import Presentation

    from export_docx import build_bundle_proposal_docx
    from export_pptx import generate_bundle_proposal
    from export_worker import sections_for
    from hcl_models import HclPlatform
    from orm_models import (CpuCatalog, DriveCatalog, DriveTypeIops, Model,
                            ModelCpuOption, RamOption, SizingSetting,
                            StorageConfig, StorageConfigDrive)
    from project_models import ExportJob
    from recommend import generate_recommendations
    from tunables import DEFAULTS

    summary = {
        "active_vms": 40, "total_vms": 44, "total_vcpus": 180,
        "total_ram_gb": 900, "used_storage_tb": 22.5, "total_storage_tb": 60.0,
        "hosts": 4, "total_host_ghz": 400.0, "peak_cpu_ghz": 120.0,
        "total_host_cores": 96, "total_host_ram_gb": 1024, "vm_iops": 0,
        "peak_ram_gb": 700, "total_vm_provisioned_memory_gb": 900,
        "datastore_used_tb": 22.5, "nic_speed_mbps": 10000,
    }

    for key, value in DEFAULTS.items():
        db.session.add(SizingSetting(key=key, value=float(value)))
    nvme = DriveCatalog(drive_type="NVMe", size_tb=7.68)
    db.session.add_all([nvme, DriveTypeIops(drive_type="NVMe", iops=75000)])
    db.session.flush()
    cpu = CpuCatalog(description="Xeon Gold 6526Y", cores=32, threads=64, ghz=2.4)
    model = Model(name="HC3450DF", status="Active", category="3XXX Core",
                  form_factor="1U", chassis="Catalog chassis", min_nodes=2,
                  cost_tier=5.0)
    db.session.add_all([cpu, model])
    db.session.flush()
    storage = StorageConfig(model_id=model.id, storage_type="nvme_only",
                            drives_per_node=4)
    db.session.add(storage)
    db.session.flush()
    db.session.add_all([
        StorageConfigDrive(storage_config_id=storage.id, drive_id=nvme.id),
        ModelCpuOption(model_id=model.id, cpu_id=cpu.id, quantity=2),
        # The sized chassis, and the one the partner quotes instead.
        HclPlatform(brand="dell", sc_model="HC3450DF",
                    server="PowerEdge R660", form_factor="1U"),
    ])
    for size in (256, 512, 1024):
        db.session.add(RamOption(model_id=model.id, size_gb=size))
    db.session.commit()

    result = generate_recommendations(summary, 4.0, sizing_mode="validated",
                                      vendor="dell")
    recs = result["recommendations"]
    assert recs, "the seeded catalog produced no Validated candidates"
    rec = recs[0]
    assert rec["vendor_chassis"] == "Dell PowerEdge R660"

    tenant = Tenant(domain="e2e.example")
    db.session.add(tenant)
    db.session.flush()
    user = User(email="e2e@e2e.example", password_hash="x",
                tenant_id=tenant.id, is_verified=True)
    db.session.add(user)
    db.session.flush()
    project = Project(name="E2E", owner_id=user.id, tenant_id=tenant.id,
                      code="PRJE")
    db.session.add(project)
    db.session.flush()
    config = Configuration(
        code="SZE", name="Site A", owner_id=user.id, tenant_id=tenant.id,
        project_id=project.id, payload={"mode": "import"},
        result_snapshot={"clusters": [{
            "name": "Prod", "summary": summary,
            "projection": result["projection"], "recommendation": rec,
            "refs": rec.get("refs") or {"mode": "appliance"},
        }], "totals": {}, "tunables": "x"})
    db.session.add(config)
    db.session.commit()
    _add_check(config)   # a passing BOM check for the R670

    job = ExportJob(user_id=user.id, project_id=project.id, fmt="pptx",
                    sizing_ids=[config.id], lang="en")
    db.session.add(job)
    db.session.commit()
    sections, skipped = sections_for(job)
    assert not skipped

    pptx = Presentation(io.BytesIO(
        generate_bundle_proposal(sections, lang="en").getvalue()))
    text = []
    for slide in pptx.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    text.append(" | ".join(c.text for c in row.cells))
    pptx_text = "\n".join(text)

    doc = Document(io.BytesIO(
        build_bundle_proposal_docx(sections, lang="en").getvalue()))
    docx_text = "\n".join([p.text for p in doc.paragraphs] +
                          [" | ".join(c.text for c in row.cells)
                           for table in doc.tables for row in table.rows])

    for name, body in (("pptx", pptx_text), ("docx", docx_text)):
        assert "PowerEdge R670" in body, f"{name} never names the quoted chassis"
        # The sized chassis must not survive anywhere in a customer document,
        # and neither must the SC model (which exports never show).
        assert "PowerEdge R660" not in body, f"{name} still names the sized chassis"
        assert "HC3450DF" not in body, f"{name} leaked the SC model"
        # the quoted per-node hardware, not the sized 32C/512GB node
        assert "6542Y" in body, f"{name} shows the sized CPU"

    # An undersized quote takes the over-100 % bar path inside a real build:
    # it must render, and the document must still name the quoted box.
    BomCheck.query.delete()
    db.session.commit()
    _add_check(config, name="Too small", fit_verdict="match",
               result=_bom_result(nodes=2, cores=8, threads=16, ram=64))
    sections, _skipped = sections_for(job)
    util = sections[0]["recommendation"]["utilization"]
    assert util["cpu"]["total"] > 100 or util["ram"]["total"] > 100
    small = Presentation(io.BytesIO(
        generate_bundle_proposal(sections, lang="en").getvalue()))
    assert any("PowerEdge R670" in shape.text_frame.text
               for slide in small.slides for shape in slide.shapes
               if shape.has_text_frame)
    assert build_bundle_proposal_docx(sections, lang="en").getvalue()


def test_put_ignores_fields_that_are_not_override_fields(client, sizing):
    client.put(f"/api/sizings/{sizing.id}/export-override",
               json={"bom": "auto", "manual": {"score": 1, "model": "HC9999",
                                               "chassis": "R670"}})
    stored = db.session.get(Configuration, sizing.id).export_override
    assert stored["manual"] == {"chassis": "R670"}


# ── the quoted card's endpoint and the renderers (2026-09-17) ───────────────

def test_preview_returns_the_quoted_card_through_the_export_code(client, sizing):
    _add_check(sizing)
    res = client.post(f"/api/sizings/{sizing.id}/export-override/preview",
                      json={"recommendation": _rec()})
    assert res.status_code == 200
    body = res.get_json()
    quoted = body["recommendation"]
    # identical to what the exports get for the same base
    expected = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    for key in ("vendor_chassis", "cores_per_node", "totals", "utilization"):
        assert quoted[key] == expected[key], key
    assert quoted["network_svg"]                      # regenerated for 5 nodes
    assert body["effective"]["chassis"] == "Dell PowerEdge R670"


def test_preview_is_empty_without_an_override_or_for_certified(client, sizing):
    res = client.post(f"/api/sizings/{sizing.id}/export-override/preview",
                      json={"recommendation": _rec()})
    assert res.get_json()["recommendation"] is None
    _add_check(sizing)
    certified = _rec(validated=False)
    res = client.post(f"/api/sizings/{sizing.id}/export-override/preview",
                      json={"recommendation": certified})
    assert res.get_json()["recommendation"] is None


def test_preview_requires_a_recommendation(client, sizing):
    res = client.post(f"/api/sizings/{sizing.id}/export-override/preview", json={})
    assert res.status_code == 400


def test_gauges_draw_an_over_capacity_block():
    """A block past 100 % renders on a rescaled axis instead of clamping; the
    image is as tall as a normal block (same rows), and the renderer does not
    fail on a bar that is nearly all overflow."""
    import io
    from PIL import Image
    from export_gauges import render_util_bars

    normal = [{"label": "CPU", "now": 35, "sized": 57, "ha": 43},
              {"label": "Storage", "now": 29, "sized": 63, "ha": 0, "snap": 11}]
    over = [{"label": "CPU", "now": 142, "sized": 300, "ha": 20},
            {"label": "Storage", "now": 29, "sized": 63, "ha": 0, "snap": 11}]
    a = Image.open(io.BytesIO(render_util_bars(normal)))
    b = Image.open(io.BytesIO(render_util_bars(over)))
    assert a.size == b.size
    # the overflow red is actually on the image
    from export_gauges import OVER
    red = tuple(int(OVER[i:i + 2], 16) for i in (1, 3, 5)) if isinstance(OVER, str) else tuple(OVER[:3])
    pixels = b.convert("RGB").getdata()
    assert any(abs(p[0] - red[0]) < 8 and abs(p[1] - red[1]) < 8 and abs(p[2] - red[2]) < 8
               for p in pixels)


def test_a_bom_check_stores_its_disk_sizes():
    """bom/fit.py keeps per-node disk sizes in the stored cluster summary, which
    is what lets the quoted card say '3 x 8 TB HDD' and re-lay the disks out."""
    import os as _os
    from bom import fit
    from bom.parsers import parse_file

    path = _os.path.join(_os.path.dirname(__file__), "fixtures", "bom",
                         "synthetic_dell_solution_de.xlsx")
    bom, _fmt = parse_file(path)
    summary = fit._cluster_summary(fit.cluster_from_nodes(fit.derive_nodes(bom.configs[0])))
    assert sorted((d["kind"], d["capacity_tb"], d["qty_per_node"]) for d in summary["drives"]) == [
        ("HDD", 8.0, 3), ("NVMe", 7.68, 1)]
    assert summary["raw_per_node_tb"] == 31.68 and summary["biggest_disk_tb"] == 8.0


def test_typed_cores_without_threads_keep_the_sized_threads_per_core(sizing):
    sizing.export_override = {"bom": "none", "manual": {"cores_per_node": 16}}
    out = eo.apply_to_rec(_rec(), eo.resolve(sizing))
    assert out["threads_per_node"] == 32          # sized 64T / 32C = 2 per core


def test_the_rationale_keeps_its_requirement_but_reports_quoted_capacity(sizing):
    """Which resource drove the sizing and what it required is the engine's
    reasoning; achieved and headroom are capacity and follow the quote - here
    negative, because a 2 x 16C quote cannot meet 84 required N-1 cores."""
    rec = _rec()
    rec["determinant"] = {"resource": "CPU", "required": 84.0, "achieved": 90.0,
                          "unit": "cores", "headroom_pct": 7.1}
    sizing.export_override = {"bom": "none", "manual": {"node_count": 2,
                                                        "cores_per_node": 16}}
    out = eo.apply_to_rec(rec, eo.resolve(sizing))
    det = out["determinant"]
    assert (det["resource"], det["required"]) == ("CPU", 84.0)
    assert det["achieved"] == 14.0                 # (16 - 2) x (2 - 1)
    assert det["headroom_pct"] == round((14 - 84) / 84 * 100, 1)
