"""The best-effort spreadsheet tier (bom/parsers/sheet_guess.py +
bom/sheet_certainty.py): a layout no strict parser knows but that is nearly
one — a DCSC export retyped with a header and the prices gone, a Dell quote
without its 'Quote number' cell, the columns in another order — is read by
column shape and trusted only when the score reaches the threshold. Anything
the read had to guess around (no quantities, a machine count taken from the
chassis line, lines it could not place) costs enough to send the file to the
agent instead, so a wrong read is never a trusted one."""
import io
import os
import sys

import pytest
from openpyxl import Workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "app"))

from bom import sheet_certainty  # noqa: E402
from bom.parsers import UnrecognizedFormat, detect_format, read_sheet_guess, sheet_guess  # noqa: E402

FIXTURES = os.path.join(ROOT, "tests", "fixtures", "bom")

LENOVO_ROWS = [
    ("7ZZ9CTO1WW", "Acme PROD : ThinkSystem SR650 V4-3yr Base Warranty", 3),
    ("Z001", 'ThinkSystem SR650 V4 12x3.5" Chassis', 3),
    ("Z002", "Intel Xeon 6517P 16C 190W 3.2GHz Processor", 6),
    ("Z003", "ThinkSystem 32GB TruDDR5 6400MHz (2Rx8) RDIMM", 48),
    ("Z004", "ThinkSystem 440-16i SAS/SATA PCIe Gen4 12Gb HBA", 3),
    ("Z005", 'ThinkSystem 2.5" U.2 VA 3.84TB Read Intensive NVMe PCIe 5.0 x4 HS SSD', 9),
    ("Z006", 'ThinkSystem 3.5" 12TB 7.2K SAS 12Gb Hot Swap 512e HDD v2', 27),
    ("Z009", "ThinkSystem Broadcom 57504 10/25GbE SFP28 4-Port OCP Ethernet Adapter", 3),
    ("Z010", "ThinkSystem 1100W 230V Titanium Hot-Swap Gen2 Power Supply v4", 6),
    ("Z014", "ThinkSystem Memory Dummy", 48),
    ("7Q01CTSAWW", "Lenovo Services", 3),
    ("QAJY", "SR650 V4 Premier Support", 3),
]

DELL_ROWS = [
    (None, "PowerEdge R760 Smart Selection - [PE_R760_TM] Production", 3),
    ("210-BDZY", "PowerEdge R760 Server", 1),
    ("461-AAIG", "Trusted Platform Module 2.0 V6", 1),
    ("404-BBED", '3.5" Chassis with up to 12 SAS/SATA Drives', 1),
    ("338-CPBV, 379-BDCO", "Intel Xeon Gold 6526Y 2.8G, 16C/32T, 20GT/s, 37.5M Cache", 2),
    ("370-AGZR", "32GB RDIMM, 5600MT/s, Dual Rank", 8),
    ("405-AAZF", "Dell HBA355i Adapter, Low Profile", 1),
    ("400-BLLG", "12TB Hard Drive SAS ISE 12Gbps 7.2K 512e 3.5in Hot-Plug", 9),
    ("540-BCXW", "Intel E810-XXV Dual Port 10/25GbE SFP28, OCP NIC 3.0", 1),
    ("450-AKYB", "Dual, Hot Plug, Power Supply (1+1) Redundant 1400W", 1),
]


def sheet(path, rows, header=("Part number", "Description", "Qty"), order=(0, 1, 2), title="Sheet1",
          extra_cols=0, blank_before=1):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for _ in range(blank_before):
        ws.append([])
    if header:
        ws.append([header[i] for i in order] + ["Price"] * extra_cols)
    for row in rows:
        ws.append([row[i] for i in order] + [12.5] * extra_cols)
    wb.save(path)
    return str(path)


def read(path):
    out = read_sheet_guess(path, os.path.basename(path))
    return out, out.certainty.score if out.certainty else None


def test_strict_ladder_still_refuses_the_retyped_layout(tmp_path):
    path = sheet(tmp_path / "retyped.xlsx", LENOVO_ROWS, extra_cols=1)
    assert detect_format(path, "retyped.xlsx") is None


def test_retyped_dcsc_with_header_is_read_and_trusted(tmp_path):
    path = sheet(tmp_path / "retyped.xlsx", LENOVO_ROWS, extra_cols=1)
    out, score = read(path)
    assert out.fmt == sheet_guess.FORMAT and out.bom.vendor == "Lenovo"
    assert score == 100 and sheet_certainty.passes(out.certainty, 85)
    (cfg,) = out.bom.configs
    assert cfg.name == "Acme PROD" and cfg.node_count == 3
    assert cfg.server_model == "ThinkSystem SR650 V4"
    by = {c.part_number: c for c in cfg.components}
    assert by["Z003"].quantity == 48 and by["Z003"].category == "memory"
    assert by["Z002"].category == "cpu" and by["Z005"].category == "storage"
    assert "Z014" not in by                       # filler dropped
    assert "QAJY" not in by and "7Q01CTSAWW" not in by   # services skipped, not unplaced
    assert out.evidence.unexplained == []


def test_columns_in_another_order_read_the_same(tmp_path):
    a = sheet(tmp_path / "a.xlsx", LENOVO_ROWS)
    b = sheet(tmp_path / "b.xlsx", LENOVO_ROWS, order=(2, 0, 1))   # Qty | Part | Description
    fa, _ = read(a)
    fb, _ = read(b)
    assert fa.bom.to_dict()["configs"] == fb.bom.to_dict()["configs"]


def test_no_header_costs_ten_points_and_still_passes(tmp_path):
    path = sheet(tmp_path / "bare.xlsx", LENOVO_ROWS, header=None)
    out, score = read(path)
    assert score == 90 and sheet_certainty.passes(out.certainty, 85)
    assert [r["code"] for r in out.certainty.reasons] == ["no_header"]


def test_dell_quote_without_its_quote_number_cell(tmp_path):
    path = sheet(tmp_path / "quote.xlsx", DELL_ROWS, header=("SKU", "Description", "Qty"))
    out, score = read(path)
    assert out.bom.vendor == "Dell" and score == 100
    (cfg,) = out.bom.configs
    # The group title opened the config and carried the machine count; the
    # server line completed it.
    assert cfg.node_count == 3 and cfg.server_model == "PowerEdge R760"
    assert cfg.name.startswith("PowerEdge R760 Smart Selection")
    by = {c.part_number: c for c in cfg.components}
    assert by["338-CPBV"].category == "cpu" and by["338-CPBV"].quantity == 2
    assert out.evidence.pattern_misses == []


def test_no_quantity_column_is_not_trusted(tmp_path):
    rows = [(p, d, None) for p, d, _q in LENOVO_ROWS]
    path = sheet(tmp_path / "noqty.xlsx", rows, header=("Part number", "Description", "Qty"))
    out, score = read(path)
    assert out.bom is not None                     # read, every line once …
    assert not sheet_certainty.passes(out.certainty, 85)   # … but not trusted
    assert "no_quantity_column" in [r["code"] for r in out.certainty.reasons]


def test_machine_count_guessed_from_the_chassis_is_not_trusted(tmp_path):
    rows = [r for r in LENOVO_ROWS[1:]]            # no CTO row: no stated count
    path = sheet(tmp_path / "nocto.xlsx", rows)
    out, score = read(path)
    (cfg,) = out.bom.configs
    assert cfg.node_count == 3                     # from the chassis line …
    codes = [r["code"] for r in out.certainty.reasons]
    assert "node_count_inferred" in codes          # … and said so
    assert score < 100


def test_unplaced_lines_cost_and_too_many_stop(tmp_path):
    blank = (None, None, None)                     # ends the services block
    noisy = LENOVO_ROWS + [blank, (None, "Intel Xeon 6517P 16C 190W 3.2GHz Processor", "see note")]
    path = sheet(tmp_path / "noisy.xlsx", noisy)
    out, score = read(path)
    assert "unexplained" in [r["code"] for r in out.certainty.reasons] and score == 88
    very = LENOVO_ROWS + [blank] + [(None, "Intel Xeon 6517P %dC Processor" % n, "n/a") for n in range(8)]
    path = sheet(tmp_path / "very.xlsx", very)
    out, score = read(path)
    assert out.certainty.hard_stop and score == 0


def test_a_sheet_without_a_parts_table_has_no_format(tmp_path):
    out, _ = read(os.path.join(FIXTURES, "not_a_bom.xlsx"))
    assert out.bom is None and out.fmt is None and out.certainty.hard_stop


def test_duplicate_summary_sheet_is_read_once(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote"
    ws.append(["Part number", "Description", "Qty"])
    for row in LENOVO_ROWS:
        ws.append(list(row))
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Part number", "Description", "Qty"])
    for row in LENOVO_ROWS:
        ws2.append(list(row))
    path = str(tmp_path / "two.xlsx")
    wb.save(path)
    out, score = read(path)
    assert len(out.bom.configs) == 1 and score == 100


def test_archive_facts_use_the_guess_only_when_trusted(tmp_path):
    from bom.facts import bom_facts
    good = sheet(tmp_path / "good.xlsx", LENOVO_ROWS)
    facts = bom_facts(good)
    assert facts["format"] == sheet_guess.FORMAT and facts["certainty"] == 100
    assert facts["configs"][0]["node_count"] == 3
    bad = sheet(tmp_path / "bad.xlsx", [(p, d, None) for p, d, _q in LENOVO_ROWS])
    with pytest.raises(UnrecognizedFormat):
        bom_facts(bad)


# ── through the route ─────────────────────────────────────────────────────────

from test_bom_routes import app, client_for, make_project  # noqa: E402,F401,F811  (its wiring)


def _post(c, pid, path):
    with open(path, "rb") as fh:
        return c.post("/api/projects/%d/bom-checks" % pid,
                      data={"file": (io.BytesIO(fh.read()), os.path.basename(path))},
                      content_type="multipart/form-data")


def test_upload_trusts_a_good_retype_and_records_the_score(app, tmp_path):
    from bom_models import BomCheck
    c = client_for(app, "retype@acme.example")
    pid = make_project(c)["id"]
    resp = _post(c, pid, sheet(tmp_path / "retyped.xlsx", LENOVO_ROWS, extra_cols=1))
    assert resp.status_code == 201, resp.get_json()
    body = resp.get_json()
    assert body["file_format"] == sheet_guess.FORMAT
    assert body["pdf"]["kind"] == "sheet" and body["pdf"]["score"] == 100
    with app.app_context():
        check = BomCheck.query.get(body["id"])
        assert check.ingest_meta["pdf"]["format"] == sheet_guess.FORMAT
        from bom.normalize import NormalizedBOM
        assert NormalizedBOM.from_dict(check.normalized).configs[0].node_count == 3


def test_upload_refuses_an_untrusted_retype_naming_why(app, tmp_path):
    c = client_for(app, "retype2@acme.example")
    pid = make_project(c)["id"]
    resp = _post(c, pid, sheet(tmp_path / "noqty.xlsx", [(p, d, None) for p, d, _q in LENOVO_ROWS]))
    # No agent configured in tests: refused, naming why the read was not trusted.
    assert resp.status_code == 400, resp.get_json()
    body = resp.get_json()
    assert "enough certainty" in body["error"]
    assert any("quantit" in d.lower() for d in body["details"])
